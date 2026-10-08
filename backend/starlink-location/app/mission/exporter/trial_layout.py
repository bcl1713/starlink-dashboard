"""Measured, fixed-font layout primitives for the editable trial deck.

Measurements use the locally bundled matplotlib DejaVu font, with a width
reserve for PowerPoint font substitution. No network/font discovery is needed.
"""

from functools import lru_cache
from pathlib import Path

import matplotlib
from PIL import ImageFont
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.util import Inches, Pt

FONT = "DejaVu Sans"
WIDTH = 13.333
HEIGHT = 7.5
INK = "263340"
NEUTRAL = "E7EAED"
POSTURE_COLORS = {
    "Nominal": "39834A",
    "Degraded": "E4B83F",
    "Limited / elevated risk": "DD8736",
    "Communications unavailable": "BD4040",
    "Posture uncertain": "E7EAED",
}
LANES = (
    "Overall communications posture",
    "Commercial Ka",
    "Starshield",
    "X-Band MILSATCOM",
    "SOF / AR restrictions",
)
LANE_HEIGHTS = (1.00, 0.44, 0.44, 0.44, 0.70)
COLUMNS = (
    "Start (ET)",
    "End (ET)",
    "Event / impact",
    "Communications remaining",
    "Overall posture / customer implication",
)
COLUMN_WIDTHS = (2.00, 2.00, 3.433, 2.70, 2.50)


@lru_cache(maxsize=16)
def _font(size, bold=False):
    suffix = "-Bold" if bold else ""
    path = Path(matplotlib.get_data_path()) / f"fonts/ttf/DejaVuSans{suffix}.ttf"
    return ImageFont.truetype(str(path), round(size * 4))


def measured_width(text, size, bold=False):
    return _font(size, bold).getlength(text) / (4 * 72)


def wrap(text, width, size=18, bold=False):
    """Explicit lines also break long source tokens; never discard characters."""
    capacity = max(0.10, width - 0.20) * 0.94
    output = []
    for paragraph in str(text).split("\n"):
        line = ""
        for word in paragraph.split(" "):
            candidate = (line + " " + word) if line else word
            if measured_width(candidate, size, bold) <= capacity:
                line = candidate
                continue
            if line:
                output.append(line)
            line = ""
            for char in word:
                if line and measured_width(line + char, size, bold) > capacity:
                    output.append(line)
                    line = ""
                line += char
        output.append(line)
    return "\n".join(output)


def line_height(size):
    return size * 1.20 / 72


def text_height(text, size=18):
    return len(text.split("\n")) * line_height(size) + 0.08


def chunks(text, width, size, height):
    lines = wrap(text, width, size).split("\n")
    count = max(1, int((height - 0.08) / line_height(size)))
    return tuple("\n".join(lines[i : i + count]) for i in range(0, len(lines), count))


def style_frame(frame, text, size=18, bold=False, color=INK):
    frame.clear()
    frame.word_wrap = False
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = MSO_ANCHOR.TOP
    frame.margin_left = frame.margin_right = Inches(0.08)
    frame.margin_top = frame.margin_bottom = Inches(0.04)
    for index, line in enumerate(text.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.space_before = paragraph.space_after = Pt(0)
        paragraph.line_spacing = Pt(size * 1.20)
        run = paragraph.add_run()
        run.text = line
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(color)


def textbox(
    slide, name, text, x, y, width, height=None, size=18, bold=False, wrapped=True
):
    rendered = wrap(text, width, size, bold) if wrapped else text
    needed = text_height(rendered, size)
    if height is not None and needed > height + 0.02:
        raise ValueError(
            f"Text exceeds fixed layout: {name} ({needed:.2f} > {height:.2f})"
        )
    shape = slide.shapes.add_textbox(
        Inches(x), Inches(y), Inches(width), Inches(height or needed)
    )
    shape.name = name
    style_frame(shape.text_frame, rendered, size, bold)
    return shape
