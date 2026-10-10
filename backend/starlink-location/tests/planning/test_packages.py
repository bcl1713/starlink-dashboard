"""Portable archives retain the entire immutable planning graph."""

import io
import zipfile
from pathlib import Path

import pytest

from app.mission import storage
from app.mission.package.__main__ import export_mission_package

from . import test_proposals, test_store
from .test_store import bind, create

service = test_store.service
prepared = test_proposals.prepared


def archive(service, mission_id):
    return export_mission_package(
        mission_id, service.store.route_manager, service.store.poi_manager
    )


def test_roundtrip_with_zero_executable_legs(service):
    from app.mission.planning.packages import commit_package, stage_package

    view, _ = create(service)
    before = service.store._load(view.mission.id)[1]
    with archive(service, view.mission.id) as handle, zipfile.ZipFile(handle) as zf:
        plan = stage_package(zf, None, service.sources)
    result = commit_package(plan, None)
    assert result.mission.id != view.mission.id
    assert result.mission.legs == []
    assert [c.leg for c in result.expected_legs] == before.expected_legs
    imported = service.store._load(result.mission.id)[1]
    assert (
        service.sources.path(imported.source_revisions[0]).read_bytes()
        == b"synthetic pdf"
    )


def test_export_missing_retained_file_is_explicit(service, tmp_path):
    view, _ = create(service)
    bind(service, tmp_path)
    _mission, manifest = service.store._load(view.mission.id)
    service.sources.path(manifest.source_revisions[0]).unlink()
    with pytest.raises((RuntimeError, ValueError, OSError)):
        archive(service, view.mission.id)


def test_collision_with_other_active_mission_remaps_complete_graph(service, tmp_path):
    from app.mission.models import MissionLeg, TransportConfig
    from app.mission.planning.packages import commit_package, stage_package

    view, _ = create(service)
    view = bind(service, tmp_path)
    mission, manifest = service.store._load(view.mission.id)
    route = manifest.expected_legs[0].route
    handle = archive(service, mission.id)
    mission.legs = [
        MissionLeg(
            id="active",
            name="Active",
            route_id=route.route_id,
            is_active=True,
            transports=TransportConfig(initial_x_satellite_id="SOUTH"),
        )
    ]
    storage.save_mission_v2(mission)
    before = storage.load_mission_v2(mission.id)
    from app.mission.exporter.snapshot import capture_export_snapshot

    before_snapshot = capture_export_snapshot(
        mission.id, service.store.route_manager, service.store.poi_manager
    )
    before_bytes = {
        str(p): p.read_bytes()
        for base in (service.sources.root, service.sources.routes_dir)
        for p in base.rglob("*")
        if p.is_file() and p.suffix != ".lock"
    }
    with handle, zipfile.ZipFile(handle) as zf:
        plan = stage_package(zf, None, service.sources)
    result = commit_package(plan, None)
    assert result.expected_legs[0].leg.route.route_id != route.route_id
    assert result.expected_legs[0].leg.route.content_hash == route.content_hash
    assert storage.load_mission_v2(mission.id) == before
    after_snapshot = capture_export_snapshot(
        mission.id, service.store.route_manager, service.store.poi_manager
    )
    assert after_snapshot.fingerprint == before_snapshot.fingerprint
    assert after_snapshot.source_payloads == before_snapshot.source_payloads
    assert all(Path(path).read_bytes() == data for path, data in before_bytes.items())
    imported_id = result.expected_legs[0].leg.route.route_id
    manager = service.store.route_manager
    cached = manager.get_route(imported_id)
    assert service.sources.resolve_profile(imported_id) == "planning_v1"
    manager._routes.pop(imported_id)
    manager._load_route_file(str(service.sources.routes_dir / f"{imported_id}.kml"))
    reloaded = manager.get_route(imported_id)
    assert reloaded.points == cached.points
    assert reloaded.timing_profile == cached.timing_profile
    assert reloaded.content_hash == cached.content_hash


@pytest.mark.parametrize("bad", ["missing", "../escape.json"])
def test_missing_file_and_invalid_archive_path_do_not_commit(service, tmp_path, bad):
    from app.mission.planning.packages import stage_package

    view, _ = create(service)
    payload = io.BytesIO()
    with archive(service, view.mission.id) as handle, zipfile.ZipFile(
        handle
    ) as original, zipfile.ZipFile(payload, "w") as target:
        for name in original.namelist():
            if bad == "missing" and name.endswith(".pdf"):
                continue
            target.writestr(name, original.read(name))
        if bad != "missing":
            target.writestr(bad, b"bad")
    before = {
        str(p): p.read_bytes()
        for p in tmp_path.rglob("*")
        if p.is_file() and p.suffix != ".lock"
    }
    with zipfile.ZipFile(payload) as zf, pytest.raises((ValueError, RuntimeError)):
        stage_package(zf, None, service.sources)
    assert before == {
        str(p): p.read_bytes()
        for p in tmp_path.rglob("*")
        if p.is_file() and p.suffix != ".lock"
    }


def test_partial_ordinals_flow_to_snapshot_excel_and_slide_identity(service):
    import json

    import openpyxl

    from app.mission.exporter.excel_utils import create_mission_summary_sheet
    from app.mission.exporter.snapshot import capture_export_snapshot
    from app.mission.models import MissionLeg, TransportConfig
    from app.mission.slide_cache.identity import leg_inputs

    view, _ = create(service)
    mission, manifest = service.store._load(view.mission.id)
    for ordinal, identity in [(3, "aaa-leg"), (1, "zzz-leg")]:
        manifest.expected_legs[ordinal - 1].installed_leg_id = identity
        mission.legs.append(
            MissionLeg(
                id=identity,
                name=identity,
                route_id="",
                transports=TransportConfig(initial_x_satellite_id="SOUTH"),
            )
        )
    service.store.persist(
        mission,
        manifest,
        files={
            storage.get_mission_leg_file_path(
                mission.id, leg.id
            ): leg.model_dump_json().encode()
            for leg in mission.legs
        },
    )
    captured = capture_export_snapshot(
        mission.id, service.store.route_manager, service.store.poi_manager
    )
    assert [leg.display_number for leg in captured.legs] == [1, 3]
    assert captured.leg_count == 3
    wb = openpyxl.Workbook()
    sheet = create_mission_summary_sheet(wb, storage.load_mission_v2(mission.id))
    assert sheet.cell(3, 2).value == 3
    assert [sheet.cell(row, 1).value for row in (7, 8)] == [1, 3]
    _, data = leg_inputs(
        captured.metadata_json, captured.source_payloads, "aaa-leg", "test"
    )
    assert json.loads(data)["number"] == 3
    assert json.loads(data)["count"] == 3
    from app.mission.exporter.customer_projection import project_briefing_leg
    from app.mission.exporter.customer_view import project_customer_leg
    from app.mission.planning.packages import commit_package, stage_package

    with archive(service, mission.id) as handle, zipfile.ZipFile(handle) as zf:
        plan = stage_package(zf, None, service.sources)
    imported = commit_package(plan, None)
    restored = capture_export_snapshot(
        imported.mission.id, service.store.route_manager, service.store.poi_manager
    )
    assert [leg.display_number for leg in restored.legs] == [1, 3]
    assert restored.leg_count == 3
    assert [
        project_customer_leg(
            leg, project_briefing_leg(leg), leg_number=leg.display_number, leg_count=3
        ).title.split(" — ")[0]
        for leg in restored.legs
    ] == ["LEG 1 OF 3", "LEG 3 OF 3"]
    before_fingerprint, _ = leg_inputs(
        captured.metadata_json, captured.source_payloads, "aaa-leg", "test"
    )
    manifest.expected_legs[0].ordinal, manifest.expected_legs[2].ordinal = 3, 1
    service.store.persist(mission, manifest)
    reordered = capture_export_snapshot(
        mission.id, service.store.route_manager, service.store.poi_manager
    )
    fingerprint, data = leg_inputs(
        reordered.metadata_json, reordered.source_payloads, "aaa-leg", "test"
    )
    assert fingerprint != before_fingerprint
    assert json.loads(data)["number"] == 1


@pytest.mark.parametrize("replace_target", [False, True])
def test_remapped_context_preserves_canonical_cost(
    prepared, monkeypatch, replace_target
):
    from app.mission.planning.evaluate import evaluate_context
    from app.mission.planning.inputs import build_inputs
    from app.mission.planning.models import ApplyProposal
    from app.mission.planning.packages import commit_package, stage_package

    from .test_proposals import proposal_api

    api, view, request = proposal_api(prepared, monkeypatch)
    proposal = api.generate(view.mission.id, "card-1", request)
    original = api.apply(
        view.mission.id,
        "card-1",
        ApplyProposal(
            expected_revision=view.revision,
            input_identity=request.input_identity,
            proposal_id=proposal.id,
        ),
    )
    service = prepared[0]
    current, retained = service.store._load(view.mission.id)
    with archive(service, view.mission.id) as handle, zipfile.ZipFile(handle) as zf:
        plan = stage_package(zf, current if replace_target else None, service.sources)
    imported = commit_package(plan, retained.revision if replace_target else None)
    if replace_target:
        assert imported.revision == retained.revision + 1
        assert imported.mission.id == current.id
    leg = imported.expected_legs[0].leg
    context = original.expected_legs[0].leg.draft.evaluation_context
    assert leg.draft.evaluation_context == context
    inputs = build_inputs(
        leg,
        leg.draft,
        service.store.route_manager,
        service.store.poi_manager,
        service.store.planning_constraints_provider(),
    )
    evaluation = evaluate_context(inputs, leg.draft, context)
    assert evaluation.outage_seconds == proposal.candidate_evaluation.outage_seconds
    assert evaluation.swap_count == proposal.candidate_evaluation.swap_count
    assert (
        evaluation.longest_gap_seconds
        == proposal.candidate_evaluation.longest_gap_seconds
    )
    _, manifest = service.store._load(imported.mission.id)
    assert (
        api._payload(imported.mission.id, manifest.proposal_refs[0]).context == context
    )


def test_staged_destination_race_and_journal_failure_rollback(
    service, tmp_path, monkeypatch
):
    from app.mission.planning import journal
    from app.mission.planning.errors import PlanningFailure
    from app.mission.planning.packages import commit_package, stage_package

    view, _ = create(service)
    view = bind(service, tmp_path)
    with archive(service, view.mission.id) as handle, zipfile.ZipFile(handle) as zf:
        plan = stage_package(zf, None, service.sources)
    baseline = {
        str(p): p.read_bytes()
        for p in tmp_path.rglob("*")
        if p.is_file() and p.suffix != ".lock"
    }
    actual = journal.atomic_write
    count = 0

    def fail_once(path, data):
        nonlocal count
        count += 1
        if count == 4:
            raise OSError("injected publication failure")
        actual(path, data)

    monkeypatch.setattr(journal, "atomic_write", fail_once)
    with pytest.raises(OSError):
        commit_package(plan, None)
    assert baseline == {
        str(p): p.read_bytes()
        for p in tmp_path.rglob("*")
        if p.is_file() and p.suffix != ".lock"
    }
    monkeypatch.setattr(journal, "atomic_write", actual)
    collision = next(Path(path) for path, data in plan.files if path.endswith(".kml"))
    collision.parent.mkdir(parents=True, exist_ok=True)
    collision.write_bytes(b"foreign race")
    with pytest.raises(PlanningFailure):
        commit_package(plan, None)
    assert collision.read_bytes() == b"foreign race"


def test_api_import_uses_complete_planning_transaction(service):
    import asyncio

    from fastapi import UploadFile
    from starlette.requests import Request

    from app.mission.routes_v2 import import_mission

    view, _ = create(service)
    handle = archive(service, view.mission.id)
    request = Request(
        {
            "type": "http",
            "method": "POST",
            "path": "/api/v2/missions/import",
            "headers": [],
        }
    )
    result = asyncio.run(
        import_mission(
            request,
            UploadFile(file=handle, filename="mission.zip"),
            service.store.route_manager,
            service.store.poi_manager,
        )
    )
    assert result["success"]
    restored = service.store.read(result["mission_id"])
    assert len(restored.expected_legs) == 3
    assert restored.mission.legs == []
    handle.close()


@pytest.mark.parametrize("retire", [False, True])
def test_partial_review_and_retired_history_roundtrip(prepared, retire):
    from app.mission.planning.models import LegHistory
    from app.mission.planning.packages import commit_package, stage_package

    from .test_reviewed_save import confirmed, review

    service, view = confirmed(prepared)
    view = service.save_reviewed(view.mission.id, "card-1", review(view))
    mission, manifest = service.store._load(view.mission.id)
    timeline = storage.load_mission_timeline(
        mission.legs[0].id, parent_mission_id=mission.id
    )
    manifest.leg_history.append(
        LegHistory(
            leg=manifest.expected_legs[0].model_copy(deep=True),
            installed_leg=mission.legs[0],
            timeline=timeline,
            source_ids=[s.id for s in manifest.source_revisions],
            revision=manifest.revision,
            reason="retirement" if retire else "revision",
        )
    )
    files = {}
    if retire:
        leg = mission.legs.pop()
        manifest.expected_legs[0].retired = True
        for card in manifest.expected_legs[1:]:
            card.ordinal -= 1
        files[storage.get_mission_leg_file_path(mission.id, leg.id)] = None
        files[storage.get_leg_timeline_path(leg.id, mission.id)] = None
    service.store.persist(mission, manifest, files=files)
    with archive(service, mission.id) as handle, zipfile.ZipFile(handle) as zf:
        plan = stage_package(zf, None, service.sources)
    imported = commit_package(plan, None)
    _, after = service.store._load(imported.mission.id)
    assert after.leg_history[0].timeline == timeline
    assert (
        after.expected_legs[0].draft.evaluation_context
        == manifest.expected_legs[0].draft.evaluation_context
    )
    assert after.expected_legs[0].review == manifest.expected_legs[0].review
    assert (
        after.leg_history[0].installed_leg.route_id
        == after.expected_legs[0].route.route_id
    )
    assert len(imported.mission.legs) == (0 if retire else 1)
    assert after.leg_history[0].source_ids != manifest.leg_history[0].source_ids
    assert not any(leg.is_active for leg in imported.mission.legs)


def test_api_rejects_missing_owned_source_before_any_mutation(service, tmp_path):
    import asyncio

    from fastapi import HTTPException, UploadFile
    from starlette.requests import Request

    from app.mission.routes_v2 import import_mission

    view, _ = create(service)
    payload = io.BytesIO()
    with archive(service, view.mission.id) as handle, zipfile.ZipFile(
        handle
    ) as original, zipfile.ZipFile(payload, "w") as output:
        for name in original.namelist():
            if not name.endswith(".pdf"):
                output.writestr(name, original.read(name))
    payload.seek(0)
    before = storage.load_mission_v2(view.mission.id)
    request = Request(
        {
            "type": "http",
            "method": "POST",
            "path": "/api/v2/missions/import",
            "headers": [],
        }
    )
    with pytest.raises(HTTPException) as raised:
        asyncio.run(
            import_mission(
                request,
                UploadFile(file=payload, filename="bad.zip"),
                service.store.route_manager,
                service.store.poi_manager,
            )
        )
    assert raised.value.status_code == 422
    assert storage.load_mission_v2(view.mission.id) == before


def test_zero_leg_pois_roundtrip_does_not_overwrite_foreign_pois(service):
    from app.mission.planning.packages import commit_package, stage_package
    from app.models.poi import POICreate

    view, _ = create(service)
    poi = service.store.poi_manager.create_poi(
        POICreate(
            name="Retained manual", latitude=1, longitude=2, mission_id=view.mission.id
        )
    )
    with archive(service, view.mission.id) as handle, zipfile.ZipFile(handle) as zf:
        plan = stage_package(zf, None, service.sources)
    result = commit_package(plan, None)
    imported = service.store.poi_manager.list_pois(mission_id=result.mission.id)
    assert len(imported) == 1
    assert imported[0].id != poi.id
    assert imported[0].name == poi.name
    assert service.store.poi_manager.list_pois(mission_id=view.mission.id) == [poi]


@pytest.mark.parametrize("change", ["owner", "profile", "hash", "path"])
def test_invalid_owned_source_graph_does_not_write(service, tmp_path, change):
    import json

    from app.mission.planning.packages import stage_package

    view, _ = create(service)
    view = bind(service, tmp_path)
    data = io.BytesIO()
    with archive(service, view.mission.id) as handle, zipfile.ZipFile(
        handle
    ) as original, zipfile.ZipFile(data, "w") as output:
        for name in original.namelist():
            content = original.read(name)
            if name.endswith(".profile.json") and change in {
                "owner",
                "profile",
                "hash",
            }:
                raw = json.loads(content)
                raw[
                    {"owner": "owner", "profile": "version", "hash": "source_hash"}[
                        change
                    ]
                ] = {"owner": "foreign", "profile": 99, "hash": "0" * 64}[change]
                content = json.dumps(raw).encode()
            if name == "mission.json" and change == "path":
                raw = json.loads(content)
                raw["metadata"]["itinerary_planning"]["source_revisions"][0][
                    "owned_relative_path"
                ] = "sources/another.pdf"
                content = json.dumps(raw).encode()
            output.writestr(name, content)
    with zipfile.ZipFile(data) as zf, pytest.raises(ValueError):
        stage_package(zf, None, service.sources)


def test_new_active_reference_after_staging_prevents_route_publication(
    service, tmp_path
):
    import json

    from app.mission.models import Mission, MissionLeg, TransportConfig
    from app.mission.planning.errors import PlanningFailure
    from app.mission.planning.packages import commit_package, stage_package

    view, _ = create(service)
    view = bind(service, tmp_path)
    with archive(service, view.mission.id) as handle, zipfile.ZipFile(handle) as zf:
        plan = stage_package(zf, None, service.sources)
    route_id = json.loads(plan.routes[0])["route_id"]
    foreign = Mission(
        id="late-owner",
        name="Late owner",
        legs=[
            MissionLeg(
                id="late-leg",
                name="Late",
                route_id=route_id,
                is_active=True,
                transports=TransportConfig(initial_x_satellite_id="SOUTH"),
            )
        ],
    )
    storage.save_mission_v2(foreign)
    with pytest.raises(PlanningFailure):
        commit_package(plan, None)
    assert not (service.sources.routes_dir / f"{route_id}.kml").exists()
    assert storage.load_mission_v2(foreign.id).legs[0].is_active


def test_import_revalidates_review_against_current_dependencies(prepared):
    from app.mission.planning.packages import commit_package, stage_package

    from .test_reviewed_save import confirmed, review

    service, view = confirmed(prepared)
    view = service.save_reviewed(view.mission.id, "card-1", review(view))
    with archive(service, view.mission.id) as handle, zipfile.ZipFile(handle) as zf:
        plan = stage_package(zf, None, service.sources)
    original = service.store.planning_constraints_provider()
    original.transition_buffer_minutes += 1
    service.store.planning_constraints_provider = lambda: original
    imported = commit_package(plan, None)
    assert imported.expected_legs[0].review_status == "needs_review"
    assert not imported.mission.legs[0].is_active
    _, manifest = service.store._load(imported.mission.id)
    assert (
        manifest.expected_legs[0].review is not None
    )  # Historical provenance survives.


def test_export_capture_blocks_reference_change_until_whole_graph_is_copied(
    service, tmp_path, monkeypatch
):
    import threading
    from concurrent.futures import ThreadPoolExecutor

    from app.mission.exporter import snapshot_inputs

    view, _ = create(service)
    view = bind(service, tmp_path)
    entered = threading.Event()
    attempted = threading.Event()
    committed = threading.Event()
    original = snapshot_inputs._read_dependencies
    calls = 0

    def dependencies(*args, **kwargs):
        nonlocal calls
        calls += 1
        if calls == 1:
            entered.set()
            assert attempted.wait(5)
            assert not committed.is_set()
        return original(*args, **kwargs)

    monkeypatch.setattr(snapshot_inputs, "_read_dependencies", dependencies)

    def writer():
        assert entered.wait(5)
        attempted.set()
        with storage.get_active_leg_lock():
            mission, manifest = service.store._load(view.mission.id)
            manifest.revision += 1
            service.store.persist(mission, manifest)
        committed.set()

    with ThreadPoolExecutor(max_workers=1) as pool:
        future = pool.submit(writer)
        metadata, _payloads, _warnings = snapshot_inputs.capture_inputs(
            view.mission.id, service.store.route_manager, service.store.poi_manager
        )
        future.result(timeout=5)
    import json

    assert (
        json.loads(metadata)["metadata"]["itinerary_planning"]["revision"]
        == view.revision
    )
    assert committed.is_set()
    assert calls == 2


def test_crash_recovery_of_import_preserves_interleaved_foreign_pois(
    service, tmp_path, monkeypatch
):
    from app.mission.planning.packages import commit_package, stage_package
    from app.models.poi import POICreate

    view, _ = create(service)
    service.store.poi_manager.create_poi(
        POICreate(
            name="Imported source", latitude=1, longitude=2, mission_id=view.mission.id
        )
    )
    with archive(service, view.mission.id) as handle, zipfile.ZipFile(handle) as zf:
        plan = stage_package(zf, None, service.sources)
    before = {
        str(p): p.read_bytes()
        for p in service.store.root.rglob("*")
        if p.is_file() and p.suffix != ".lock"
    }

    class Crash(BaseException):
        pass

    original = service.store.journal.commit

    def crash_commit(files, **kwargs):
        def crash(index, path):
            if path == service.store.poi_manager.pois_file.resolve():
                import json

                from filelock import FileLock

                from app.mission.planning.journal import atomic_write, json_bytes
                from app.models.poi import POI

                # A separately persisted foreign update must survive recovery.
                with FileLock(str(path) + ".lock"):
                    records = json.loads(path.read_bytes())
                    poi = POI(
                        id="foreign-during-import",
                        name="Foreign during import",
                        latitude=3,
                        longitude=4,
                        mission_id="foreign",
                    )
                    records["pois"][poi.id] = poi.model_dump(mode="json")
                    atomic_write(path, json_bytes(records))
                raise Crash()

        return original(files, **kwargs, after_write=crash)

    monkeypatch.setattr(service.store.journal, "commit", crash_commit)
    with pytest.raises(Crash):
        commit_package(plan, None)
    assert list(service.store.journal.directory.glob("*.json"))
    service.store.recover()
    assert not list(service.store.journal.directory.glob("*.json"))
    assert before == {
        str(p): p.read_bytes()
        for p in service.store.root.rglob("*")
        if p.is_file() and p.suffix != ".lock"
    }
    assert (
        service.store.poi_manager.list_pois(mission_id="foreign")[0].name
        == "Foreign during import"
    )
    assert len(service.store.poi_manager.list_pois(mission_id=view.mission.id)) == 1


def test_legacy_package_imports_new_satellite_without_overwriting_existing(service):
    import json

    from app.mission.models import Mission
    from app.mission.planning.packages import commit_package, stage_package
    from app.models.poi import POI, POICreate

    existing = service.store.poi_manager.create_poi(
        POICreate(
            name="Existing X", latitude=0, longitude=30, icon="X", category="satellite"
        )
    )
    mission = Mission(id="legacy-satellite", name="Legacy")
    package = io.BytesIO()
    with zipfile.ZipFile(package, "w") as zf:
        zf.writestr("mission.json", mission.model_dump_json())
        zf.writestr(
            "pois/satellites.json",
            json.dumps(
                {
                    "pois": [
                        POI(
                            id="archived-existing",
                            name="Existing X",
                            latitude=0,
                            longitude=99,
                            icon="X",
                            category="satellite",
                        ).model_dump(mode="json"),
                        POI(
                            id="archived-new",
                            name="New X",
                            latitude=0,
                            longitude=40,
                            icon="X",
                            category="satellite",
                        ).model_dump(mode="json"),
                    ]
                }
            ),
        )
    with zipfile.ZipFile(package) as zf:
        plan = stage_package(zf, None, service.sources)
    commit_package(plan, None)
    assert service.store.poi_manager.find_global_poi_by_name("Existing X") == existing
    assert service.store.poi_manager.find_global_poi_by_name("New X").longitude == 40


def test_legacy_package_missing_referenced_route_is_rejected(service):
    from app.mission.models import Mission, MissionLeg, TransportConfig
    from app.mission.planning.packages import stage_package

    mission = Mission(
        id="legacy-missing",
        name="Missing",
        legs=[
            MissionLeg(
                id="one",
                name="One",
                route_id="absent",
                transports=TransportConfig(initial_x_satellite_id="SOUTH"),
            )
        ],
    )
    data = io.BytesIO()
    with zipfile.ZipFile(data, "w") as zf:
        zf.writestr("mission.json", mission.model_dump_json())
    with zipfile.ZipFile(data) as zf, pytest.raises(ValueError, match="missing"):
        stage_package(zf, None, service.sources)
    assert storage.load_mission_v2(mission.id) is None


def test_snapshot_export_rejects_missing_captured_owned_bytes(service):
    from dataclasses import replace

    from app.mission.exporter.snapshot import capture_export_snapshot

    view, _ = create(service)
    captured = capture_export_snapshot(
        view.mission.id, service.store.route_manager, service.store.poi_manager
    )
    incomplete = replace(
        captured,
        source_payloads=tuple(
            p for p in captured.source_payloads if not p.name.endswith(".pdf")
        ),
    )
    with pytest.raises((RuntimeError, ValueError)):
        export_mission_package(
            view.mission.id,
            service.store.route_manager,
            service.store.poi_manager,
            snapshot=incomplete,
        )


@pytest.mark.parametrize("missing", ["source", "history", "work"])
def test_target_import_cannot_discard_retained_work(service, tmp_path, missing):
    from app.mission.planning.models import LegHistory
    from app.mission.planning.packages import stage_package

    view, _ = create(service)
    view = bind(service, tmp_path)
    archive_before = archive(service, view.mission.id)
    mission, manifest = service.store._load(view.mission.id)
    if missing == "source":
        source = service.sources.stage(b"later source", "pdf", mission.id, "later.pdf")
        accepted, files = service.sources.acceptance_files(
            source, mission.id, b"later source"
        )
        manifest.source_revisions.append(accepted)
    elif missing == "history":
        manifest.leg_history.append(
            LegHistory(
                leg=manifest.expected_legs[0],
                revision=manifest.revision,
                reason="revision",
            )
        )
        files = {}
    else:
        manifest.expected_legs[0].draft.starshield_enabled = False
        files = {}
    service.store.persist(mission, manifest, files=files)
    current = storage.load_mission_v2(mission.id)
    with archive_before, zipfile.ZipFile(archive_before) as zf, pytest.raises(
        ValueError
    ):
        stage_package(zf, current, service.sources)
    assert storage.load_mission_v2(mission.id) == current


def test_explicit_target_import_preserves_ids_and_advances_revision(service, tmp_path):
    from app.mission.planning.packages import commit_package, stage_package

    view, _ = create(service)
    view = bind(service, tmp_path)
    current, before = service.store._load(view.mission.id)
    with archive(service, view.mission.id) as handle, zipfile.ZipFile(handle) as zf:
        plan = stage_package(zf, current, service.sources)
    result = commit_package(plan, before.revision)
    assert result.revision == before.revision + 1
    assert result.expected_legs[0].leg.route == before.expected_legs[0].route
    assert result.mission.id == current.id


def test_explicit_target_rejects_missing_proposal_payload_and_stale_cas(
    prepared, monkeypatch
):
    from app.mission.planning.errors import PlanningFailure
    from app.mission.planning.packages import commit_package, stage_package

    from .test_proposals import proposal_api

    api, view, request = proposal_api(prepared, monkeypatch)
    service = prepared[0]
    old_archive = archive(service, view.mission.id)
    api.generate(view.mission.id, "card-1", request)
    current, retained = service.store._load(view.mission.id)
    assert retained.proposal_refs
    before = {
        str(p): p.read_bytes()
        for p in service.store.root.rglob("*")
        if p.is_file() and p.suffix != ".lock"
    }
    with old_archive, zipfile.ZipFile(old_archive) as zf, pytest.raises(
        ValueError, match="retained"
    ):
        stage_package(zf, current, service.sources)
    with archive(service, view.mission.id) as handle, zipfile.ZipFile(handle) as zf:
        plan = stage_package(zf, current, service.sources)
    with pytest.raises(PlanningFailure):
        commit_package(plan, retained.revision - 1)
    assert before == {
        str(p): p.read_bytes()
        for p in service.store.root.rglob("*")
        if p.is_file() and p.suffix != ".lock"
    }


def test_commit_rejects_parent_directory_alias_without_foreign_writes(
    service, tmp_path
):
    from app.mission.planning.errors import PlanningFailure
    from app.mission.planning.packages import commit_package, stage_package

    view, _ = create(service)
    with archive(service, view.mission.id) as handle, zipfile.ZipFile(handle) as zf:
        plan = stage_package(zf, None, service.sources)
    import json

    destination = service.store.root / json.loads(plan.mission_json)["id"]
    foreign = service.store.root / "foreign-directory"
    foreign.mkdir()
    destination.symlink_to(foreign, target_is_directory=True)
    with pytest.raises((ValueError, PlanningFailure)):
        commit_package(plan, None)
    assert list(foreign.iterdir()) == []


def test_route_replacement_roundtrip_preserves_old_anchors_locks_pending_and_baseline(
    prepared, tmp_path
):
    from app.mission.planning.match import _candidates
    from app.mission.planning.models import AnchoredSwap, PlanningLock, SaveDraft
    from app.mission.planning.packages import commit_package, stage_package

    from .test_reviewed_save import confirmed, review
    from .test_revisions import accept, replacement

    service, view = confirmed(prepared)
    leg = view.expected_legs[0].leg
    route = service.store.route_manager.get_route(leg.route.route_id)
    anchor = _candidates(route.points[1].expected_arrival_time, "second", route)[0]
    draft = leg.draft.model_copy(deep=True)
    draft.swaps = [AnchoredSwap(id="manual", target_satellite_id="WEST", anchor=anchor)]
    draft.locks = [
        PlanningLock(
            id="lock", swap_id="manual", target_satellite_id="WEST", anchor=anchor
        )
    ]
    view = service.store.save_draft(
        view.mission.id, leg.id, SaveDraft(expected_revision=view.revision, draft=draft)
    )
    view = service.save_reviewed(view.mission.id, leg.id, review(view))
    accept(service, view, replacement(service, view, tmp_path))
    original, before = service.store._load(view.mission.id)
    from app.mission.models import AARWindow, XTransition

    before.expected_legs[0].draft.unresolved_x_transitions = [
        XTransition(
            id="pending-x",
            latitude=anchor.latitude,
            longitude=anchor.longitude,
            target_satellite_id="WEST",
            anchor=anchor,
        )
    ]
    before.expected_legs[0].draft.unresolved_aar_windows = [
        AARWindow(
            id="pending-ar",
            start_waypoint_name=route.waypoints[0].name,
            end_waypoint_name=route.waypoints[-1].name,
            start_anchor=anchor,
        )
    ]
    service.store.persist(original, before)
    assert (
        before.expected_legs[0].draft.locks[0].anchor.route_id
        != before.expected_legs[0].route.route_id
    )
    with archive(service, view.mission.id) as handle, zipfile.ZipFile(handle) as zf:
        plan = stage_package(zf, None, service.sources)
    imported = commit_package(plan, None)
    restored, after = service.store._load(imported.mission.id)
    remaps = dict(plan.remaps)
    assert after.itinerary_baseline == before.itinerary_baseline
    assert after.expected_legs[0].draft.locks[0].anchor == anchor.model_copy(
        update={"route_id": remaps[anchor.route_id]}
    )
    assert after.expected_legs[0].draft.swaps[0].id == "manual"
    expected_pending = (
        before.expected_legs[0].draft.unresolved_x_transitions[0].model_copy(deep=True)
    )
    expected_pending.anchor.route_id = remaps[anchor.route_id]
    assert after.expected_legs[0].draft.unresolved_x_transitions == [expected_pending]
    assert (
        after.expected_legs[0].draft.unresolved_aar_windows[0].start_anchor
        == expected_pending.anchor
    )
    assert (
        after.expected_legs[0].draft.unresolved_aar_windows[0].override_start_time
        is None
    )
    assert len(after.route_history) == len(before.route_history)
    assert {source.content_hash for source in after.source_revisions} == {
        source.content_hash for source in before.source_revisions
    }
    assert restored.legs[0].route_id == remaps[original.legs[0].route_id]
    assert not restored.legs[0].is_active
    assert imported.expected_legs[0].leg.review is None


@pytest.mark.parametrize("matching", [False, True])
def test_api_import_verifies_manager_context_and_preserves_constraints(
    service, tmp_path, monkeypatch, matching
):
    import asyncio
    from types import SimpleNamespace

    from fastapi import UploadFile
    from starlette.requests import Request

    from app.mission.planning import packages
    from app.mission.planning.store import PlanningStore
    from app.mission.routes_v2 import import_mission
    from app.satellites.rules import ConstraintConfig
    from app.services.poi_manager import POIManager
    from app.services.route_manager import RouteManager

    view, _ = create(service)
    root = service.store.root
    if not matching:
        monkeypatch.setattr(storage, "MISSIONS_DIR", tmp_path / "other-missions")
    configured = (
        service.store
        if matching
        else PlanningStore(
            tmp_path / "other-missions",
            RouteManager(tmp_path / "other-routes"),
            POIManager(tmp_path / "other-pois.json"),
        )
    )
    monkeypatch.setattr(storage, "MISSIONS_DIR", root)

    def provider():
        return ConstraintConfig()

    configured.planning_constraints_provider = provider
    original = packages.stage_package
    observed = []

    def checked(zf, target, sources):
        observed.append(sources._store)
        assert sources._store.root == root
        assert sources._store.route_manager is service.store.route_manager
        assert sources._store.poi_manager is service.store.poi_manager
        assert sources._store.planning_constraints_provider is provider
        return original(zf, target, sources)

    monkeypatch.setattr(packages, "stage_package", checked)
    request = Request(
        {
            "type": "http",
            "method": "POST",
            "path": "/api/v2/missions/import",
            "headers": [],
            "app": SimpleNamespace(
                state=SimpleNamespace(
                    planning_service=SimpleNamespace(store=configured)
                )
            ),
        }
    )
    with archive(service, view.mission.id) as handle:
        result = asyncio.run(
            import_mission(
                request,
                UploadFile(file=handle, filename="mission.zip"),
                service.store.route_manager,
                service.store.poi_manager,
            )
        )
    assert result["success"]
    assert (observed[0] is configured) is matching


def test_package_source_context_cannot_be_unbound_or_rebound(service):
    from types import SimpleNamespace

    from app.mission.planning.packages import stage_package
    from app.mission.planning.sources import SourceStore

    view, _ = create(service)
    unbound = SourceStore(service.store.root, service.sources.routes_dir)
    with archive(service, view.mission.id) as handle, zipfile.ZipFile(
        handle
    ) as zf, pytest.raises(ValueError, match="bound"):
        stage_package(zf, None, unbound)
    with pytest.raises(ValueError, match="already has a commit context"):
        service.sources.bind_store(
            SimpleNamespace(
                sources=service.sources,
                root=service.store.root,
                route_manager=service.store.route_manager,
            )
        )
