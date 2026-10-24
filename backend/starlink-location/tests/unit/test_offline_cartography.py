"""Real cartographic rendering must work with no downloaded Natural Earth data."""

import io

import cartopy.feature as cfeature
import pytest
from cartopy.io import Downloader
from PIL import Image

from app.mission.exporter import __main__ as exporter
from app.mission.models import MissionLegTimeline

pytest_plugins = ["tests.unit.test_export_snapshot"]


def test_offline_routes_render_distinct_repeatable_valid_pngs(
    offline_cartography, export_inputs
):
    mission, routes, pois, _ = export_inputs
    timeline = MissionLegTimeline(mission_leg_id=mission.legs[0].id, segments=[])
    route_b = routes._routes["r"].model_copy(deep=True)
    for point in route_b.points:
        point.longitude += 80 + point.sequence * 8
        point.latitude += 30
    routes._routes["b"] = route_b
    leg_b = mission.legs[0].model_copy(update={"id": "b", "route_id": "b"}, deep=True)

    def render(leg):
        return exporter._generate_route_map(
            timeline, leg, route_manager=routes, poi_manager=pois
        )

    assert not list(offline_cartography["cache"].rglob("*.shp"))
    first = render(mission.legs[0])
    second = render(leg_b)
    assert first != second
    assert render(mission.legs[0]) == first
    for content in (first, second):
        with Image.open(io.BytesIO(content)) as image:
            assert image.format == "PNG"
            assert image.size == (3840, 2160)
            image.load()
            assert any(low < high for low, high in image.getextrema())
    assert {name for _category, name in offline_cartography["requested"]} == {
        "coastline",
        "admin_0_boundary_lines_land",
        "land",
        "ocean",
    }
    assert not list(offline_cartography["cache"].iterdir())


def test_offline_geometry_is_nonempty_valid_and_rejects_unknown(offline_cartography):
    for feature in (
        cfeature.COASTLINE,
        cfeature.BORDERS,
        cfeature.LAND,
        cfeature.OCEAN,
    ):
        geometries = tuple(feature.geometries())
        assert geometries
        assert all(not item.is_empty and item.is_valid for item in geometries)
    with pytest.raises(AssertionError, match="Unsupported Natural Earth"):
        tuple(cfeature.NaturalEarthFeature("physical", "lakes", "110m").geometries())


def test_cartography_download_guard_is_active(offline_cartography):
    with pytest.raises(AssertionError, match="External cartography download"):
        Downloader("https://example.invalid/map.zip", "map.zip")._urlopen(
            "https://example.invalid/map.zip"
        )


def test_fixture_restores_cartography_patches(tmp_path):
    import cartopy

    from tests.fixtures.cartography import offline_cartography as fixture

    original_geometries = cfeature.NaturalEarthFeature.geometries
    original_urlopen = Downloader._urlopen
    original_config = dict(cartopy.config)
    with pytest.MonkeyPatch.context() as patch:
        fixture.__wrapped__(patch, tmp_path)
        assert cfeature.NaturalEarthFeature.geometries is not original_geometries
        assert Downloader._urlopen is not original_urlopen
    assert cfeature.NaturalEarthFeature.geometries is original_geometries
    assert Downloader._urlopen is original_urlopen
    assert dict(cartopy.config) == original_config


def test_real_pptx_contains_route_png_and_mission_text(
    offline_cartography, export_inputs
):
    import zipfile

    from pptx import Presentation

    mission, routes, pois, _ = export_inputs
    leg = mission.legs[0]
    content = exporter.generate_pptx_export(
        MissionLegTimeline(mission_leg_id=leg.id, segments=[]),
        leg,
        route_manager=routes,
        poi_manager=pois,
    )
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        media = [name for name in archive.namelist() if name.startswith("ppt/media/")]
        assert media
        assert any(
            archive.read(name).startswith(b"\x89PNG\r\n\x1a\n") for name in media
        )
    presentation = Presentation(io.BytesIO(content))
    text = "\n".join(
        shape.text
        for slide in presentation.slides
        for shape in slide.shapes
        if shape.has_text_frame
    )
    assert leg.name in text
    assert "Route Map" in text
    assert offline_cartography["requested"]
