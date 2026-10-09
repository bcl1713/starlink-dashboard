"""Export must freeze one revision without publishing preparation results."""

import io
import json
from dataclasses import FrozenInstanceError
from datetime import datetime, timedelta, timezone
from itertools import pairwise
from unittest.mock import Mock

import pytest
from pptx import Presentation

from app.mission import storage, timeline_preparation
from app.mission.models import (
    AARWindow,
    ManualAARTrack,
    ManualAARTrackPoint,
    ManualRouteSplice,
    Mission,
    MissionLeg,
    TransportConfig,
)
from app.models.poi import POI
from app.models.route import ParsedRoute, RouteMetadata, RoutePoint, RouteTimingProfile
from app.satellites.catalog import Satellite, SatelliteCatalog
from app.services.poi_manager import POIManager
from app.services.route_manager import RouteManager

BASE = datetime(2026, 10, 7, 8, tzinfo=timezone.utc)


@pytest.fixture
def export_inputs(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(storage, "MISSIONS_DIR", tmp_path / "missions")
    route = ParsedRoute(
        metadata=RouteMetadata(name="Captured route", file_path="r.kml", point_count=2),
        points=[
            RoutePoint(
                latitude=lat, longitude=0, sequence=i, expected_arrival_time=time
            )
            for i, lat, time in [(0, 0, BASE), (1, 1, BASE + timedelta(hours=1))]
        ],
        timing_profile=RouteTimingProfile(
            departure_time=BASE,
            arrival_time=BASE + timedelta(hours=1),
            has_timing_data=True,
        ),
    )
    routes = RouteManager(tmp_path / "routes")
    routes._routes["r"] = route
    (routes.routes_dir / "r.kml").write_bytes(b"<kml>captured</kml>")
    pois = POIManager(tmp_path / "pois.json")
    pois._pois["p"] = POI(
        id="p",
        name="Captured POI",
        latitude=0,
        longitude=0,
        mission_id="m",
        route_id="r",
        category="mission-event",
        created_at=BASE,
        updated_at=BASE,
    )
    catalog = SatelliteCatalog()
    catalog.add_satellite(Satellite("X-test", "X", longitude=0))
    monkeypatch.setattr("app.satellites.catalog._catalog", catalog)
    mission = Mission(
        id="m",
        name="Captured mission",
        description="Captured description",
        metadata={"revision": "1"},
        legs=[
            MissionLeg(
                id="l",
                name="Captured leg",
                route_id="r",
                transports=TransportConfig(
                    initial_x_satellite_id="X-test",
                    aar_windows=[
                        AARWindow(
                            id="ar",
                            start_waypoint_name="unused",
                            end_waypoint_name="unused",
                            override_start_time=BASE + timedelta(minutes=20),
                            override_end_time=BASE + timedelta(minutes=40),
                        )
                    ],
                ),
            )
        ],
    )
    storage.save_mission_v2(mission)
    return mission, routes, pois, catalog


def _slide_text(content):
    return "\n".join(
        shape.text
        for slide in Presentation(io.BytesIO(content)).slides
        for shape in slide.shapes
        if shape.has_text_frame
    )


def test_enabled_package_builders_never_reread_after_capture(
    export_inputs, monkeypatch
):
    import inspect
    import zipfile

    from app.mission.exporter.customer_document import build_customer_mission_document
    from app.mission.exporter.snapshot import capture_export_snapshot
    from app.mission.package import __main__ as package

    mission, routes, pois, catalog = export_inputs
    captured = capture_export_snapshot("m", routes, pois)
    assert (
        "snapshot" in inspect.signature(package.export_mission_package).parameters
    ), "Snapshot package injection absent"
    mission.name = "Later mission"
    storage.save_mission_v2(mission)
    (routes.routes_dir / "r.kml").write_bytes(b"later KML")
    routes._routes.clear()
    pois._pois.clear()
    catalog.satellites.clear()

    def forbidden(*args, **kwargs):
        raise AssertionError("Live dependency reread after capture")

    for target in (
        "app.mission.storage.load_mission_v2",
        "app.mission.storage.load_mission_timeline",
        "app.mission.package.__main__.load_mission_v2",
        "app.mission.package.__main__.load_mission_timeline",
        "app.mission.package.__main__.build_mission_timeline",
        "app.mission.exporter.__main__.get_cached_ground_entry_point",
        "app.mission.exporter.snapshot.prepare_mission_timeline",
    ):
        monkeypatch.setattr(target, forbidden)
    monkeypatch.setattr(routes, "get_route", forbidden)
    monkeypatch.setattr(pois, "list_pois", forbidden)
    data_root = routes.routes_dir.parent
    before = {
        str(p.relative_to(data_root)): p.read_bytes()
        for p in data_root.rglob("*")
        if p.is_file()
    }
    with package.export_mission_package(
        "m", routes, pois, snapshot=captured
    ) as stream, zipfile.ZipFile(stream) as archive:
        assert json.loads(archive.read("mission.json"))["name"] == "Captured mission"
        assert archive.read("routes/r.kml") == b"<kml>captured</kml>"
        assert b"Captured POI" in archive.read("pois/l-pois.json")
        assert not any(name.endswith(".pptx") for name in archive.namelist())
        assert b"1970" not in archive.read("exports/legs/l/timeline.csv")
    assert (
        build_customer_mission_document(captured)["snapshotFingerprint"]
        == captured.fingerprint
    )
    assert {
        str(p.relative_to(data_root)): p.read_bytes()
        for p in data_root.rglob("*")
        if p.is_file()
    } == before


@pytest.mark.parametrize(
    "case",
    [
        "normal",
        "adjusted",
        "shared-route",
        "cached",
        "missing-timeline",
        "ar",
        "splice",
    ],
)
def test_snapshot_package_matches_fixed_clock_legacy(export_inputs, monkeypatch, case):
    import inspect
    import zipfile

    from app.mission.exporter.snapshot import capture_export_snapshot
    from app.mission.package import __main__ as package

    mission, routes, pois, _ = export_inputs
    assert (
        "snapshot" in inspect.signature(package.export_mission_package).parameters
    ), "Snapshot package injection absent"
    if case != "ar":
        mission.legs[0].transports.aar_windows = []
    if case == "splice":
        mission.legs[0].transports.manual_aar_tracks = [
            ManualAARTrack(
                id="track",
                name="Track",
                points=[
                    ManualAARTrackPoint(latitude=0.3, longitude=0.01),
                    ManualAARTrackPoint(latitude=0.7, longitude=0.01),
                ],
            )
        ]
        mission.legs[0].transports.manual_route_splice = ManualRouteSplice(
            enabled_track_id="track", speed_knots=100
        )
    if case == "adjusted":
        mission.legs[0].adjusted_departure_time = BASE + timedelta(hours=2)
    if case == "shared-route":
        second = mission.legs[0].model_copy(deep=True)
        second.id = "l2"
        second.name = "Second captured leg"
        mission.legs.append(second)
    storage.save_mission_v2(mission)
    if case in ("cached", "missing-timeline"):
        if case == "cached":
            storage.save_mission_timeline(
                "l",
                timeline_preparation.prepare_mission_timeline(
                    mission.legs[0], routes, pois
                ).timeline,
                parent_mission_id="m",
            )
        failing = Mock(side_effect=RuntimeError("Forced cached fallback"))
        monkeypatch.setattr(
            "app.mission.exporter.snapshot.prepare_mission_timeline", failing
        )
        monkeypatch.setattr(package, "build_mission_timeline", failing)

    class DatetimeType(type):
        def __instancecheck__(cls, value):
            return isinstance(value, datetime)

    class FixedDatetime(datetime, metaclass=DatetimeType):
        @classmethod
        def now(cls, tz=None):
            return BASE

    monkeypatch.setattr(package, "datetime", FixedDatetime)
    captured = capture_export_snapshot("m", routes, pois)
    with package.export_mission_package("m", routes, pois) as stream:
        baseline = stream.getvalue()
    with package.export_mission_package("m", routes, pois, snapshot=captured) as stream:
        frozen = stream.getvalue()
    with zipfile.ZipFile(io.BytesIO(baseline)) as a, zipfile.ZipFile(
        io.BytesIO(frozen)
    ) as b:
        assert a.namelist() == b.namelist()
        for name in a.namelist():
            if name.endswith(".pptx"):
                with zipfile.ZipFile(io.BytesIO(a.read(name))) as x, zipfile.ZipFile(
                    io.BytesIO(b.read(name))
                ) as y:
                    assert x.namelist() == y.namelist()
                    for member in x.namelist():
                        assert x.read(member) == y.read(member), (case, name, member)
            else:
                assert a.read(name) == b.read(name), (case, name)


def test_snapshot_isolated_from_later_save_and_builder_mutation(
    export_inputs, monkeypatch
):
    from app.mission.exporter.snapshot import capture_export_snapshot
    from app.mission.exporter.snapshot_views import SnapshotViews

    mission, routes, pois, _ = export_inputs
    prepare = Mock(wraps=timeline_preparation.prepare_mission_timeline)
    monkeypatch.setattr(
        "app.mission.exporter.snapshot.prepare_mission_timeline", prepare
    )
    snapshot = capture_export_snapshot("m", routes, pois)
    assert prepare.call_count == 1
    mission.name = "Later mission"
    mission.metadata["revision"] = "2"
    mission.legs[0].adjusted_departure_time = BASE + timedelta(hours=2)
    storage.save_mission_v2(mission)
    routes._routes["r"].points[0].latitude = 40
    (routes.routes_dir / "r.kml").write_bytes(b"<kml>later</kml>")
    pois._pois["p"].name = "Later POI"
    views = SnapshotViews(snapshot)
    views.mission().name = "Builder mutation"
    views.timeline("l").segments.clear()
    views.route_manager.get_route("r").points.clear()
    views.poi_manager.list_pois()[0].name = "Builder mutation"
    assert views.mission().name == "Captured mission"
    assert views.timeline("l").segments
    assert views.route_manager.get_route("r").points[0].latitude == 0
    assert views.poi_manager.list_pois()[0].name == "Captured POI"
    with pytest.raises((FrozenInstanceError, AttributeError)):
        snapshot.fingerprint = "mutation"
    later = capture_export_snapshot("m", routes, pois)
    assert later.fingerprint != snapshot.fingerprint
    assert SnapshotViews(later).mission().name == "Later mission"
    assert SnapshotViews(later).timeline("l").segments[
        0
    ].start_time == BASE + timedelta(hours=2)


def test_capture_uses_prepare_without_publication(export_inputs, monkeypatch):
    from app.mission.exporter.snapshot import capture_export_snapshot

    mission, routes, pois, catalog = export_inputs
    original_route = routes.get_route("r").model_dump_json()
    original_pois = [p.model_dump_json() for p in pois.list_pois()]
    before = storage.load_mission_v2("m").model_dump_json()

    def forbidden(*args, **kwargs):
        pytest.fail("Export published or wrote storage")

    for name in ["save_mission_v2", "save_mission_timeline"]:
        monkeypatch.setattr(storage, name, forbidden)
    for name in [
        "create_poi",
        "delete_leg_pois",
        "delete_scoped_pois_by_names",
        "_save_pois",
    ]:
        monkeypatch.setattr(pois, name, forbidden)
    monkeypatch.setattr(timeline_preparation, "publish_mission_pois", forbidden)
    original = timeline_preparation.prepare_mission_timeline

    def prepare(*args, **kwargs):
        assert not storage.get_active_leg_lock().is_locked
        assert not storage.get_mission_lock("m").is_locked
        # Live changes during computation must not influence frozen dependencies.
        catalog.satellites["X-test"].longitude = 180
        return original(*args, **kwargs)

    monkeypatch.setattr(
        "app.mission.exporter.snapshot.prepare_mission_timeline", prepare
    )
    snapshot = capture_export_snapshot("m", routes, pois)
    leg = snapshot.legs[0]
    assert leg.preparation_origin == "rebuilt"
    assert leg.utc_bounds == (BASE, BASE + timedelta(hours=1))
    assert leg.source_records
    assert not any(
        "Elevation below" in record.decode() for record in leg.source_records
    )
    assert any(
        json.loads(record)["source_id"] == "ar" for record in leg.resolved_restrictions
    )
    assert storage.load_mission_v2("m").model_dump_json() == before
    assert routes.get_route("r").model_dump_json() == original_route
    assert [p.model_dump_json() for p in pois.list_pois()] == original_pois
    assert mission.legs[0].adjusted_departure_time is None


@pytest.mark.parametrize("cached", [True, False])
@pytest.mark.parametrize("missing_route", [True, False])
def test_snapshot_rebuild_cache_missing(
    export_inputs, monkeypatch, cached, missing_route
):
    from app.mission.exporter.snapshot import capture_export_snapshot
    from app.mission.exporter.snapshot_views import SnapshotViews

    mission, routes, pois, _ = export_inputs
    if cached:
        old = timeline_preparation.prepare_mission_timeline(
            storage.load_mission_v2("m").legs[0], routes, pois
        ).timeline
        storage.save_mission_timeline("l", old, parent_mission_id="m")
    mission.legs[0].adjusted_departure_time = BASE + timedelta(hours=2)
    storage.save_mission_v2(mission)
    if missing_route:
        routes._routes.clear()
    monkeypatch.setattr(
        "app.mission.exporter.snapshot.prepare_mission_timeline",
        Mock(side_effect=RuntimeError("rebuild failure")),
    )
    snapshot = capture_export_snapshot("m", routes, pois)
    leg = snapshot.legs[0]
    assert leg.preparation_origin == ("cached" if cached else "missing")
    assert leg.utc_bounds == (
        None
        if missing_route
        else (BASE + timedelta(hours=2), BASE + timedelta(hours=3))
    )
    assert any(
        ("cached" if cached else "missing") in warning.lower()
        for warning in leg.warnings
    )
    if cached:
        assert SnapshotViews(snapshot).timeline("l").segments[0].start_time == BASE
        assert (
            leg.source_records == ()
        )  # No invented canonical events from normalized cache.
    else:
        assert leg.timeline_json is None


@pytest.mark.parametrize("available", [True, False])
def test_snapshot_effective_splice_and_restriction_identity(export_inputs, available):
    from app.mission.exporter.snapshot import capture_export_snapshot
    from app.mission.exporter.snapshot_views import SnapshotViews

    mission, routes, pois, _ = export_inputs
    mission.legs[0].transports.manual_aar_tracks = [
        ManualAARTrack(
            id="track",
            name="Track",
            points=[
                ManualAARTrackPoint(latitude=0.3 if available else 70, longitude=0.01),
                ManualAARTrackPoint(latitude=0.7 if available else 71, longitude=0.01),
            ],
        )
    ]
    mission.legs[0].transports.manual_route_splice = ManualRouteSplice(
        enabled_track_id="track", speed_knots=100
    )
    storage.save_mission_v2(mission)
    snapshot = capture_export_snapshot("m", routes, pois)
    effective = json.loads(snapshot.legs[0].effective_route_json)
    assert (len(effective["points"]) > 2) is available
    restrictions = [
        json.loads(record) for record in snapshot.legs[0].resolved_restrictions
    ]
    assert any(r["source_id"] == "track" for r in restrictions) is available
    assert SnapshotViews(snapshot).route_manager.get_route("r").points[0].latitude == 0


def test_capture_retries_changed_dependencies_once(export_inputs, monkeypatch):
    from app.mission.exporter.snapshot import capture_export_snapshot
    from app.mission.exporter.snapshot_views import SnapshotViews

    _, routes, pois, _ = export_inputs
    original = routes.get_route
    calls = 0

    def changing(route_id):
        nonlocal calls
        calls += 1
        if calls == 2:
            routes._routes["r"].metadata.name = "New revision"
        return original(route_id)

    monkeypatch.setattr(routes, "get_route", changing)
    snapshot = capture_export_snapshot("m", routes, pois)
    assert (
        SnapshotViews(snapshot).route_manager.get_route("r").metadata.name
        == "New revision"
    )
    assert any("retry" in warning.lower() for warning in snapshot.warnings)


def test_capture_rejects_continuously_changing_inputs(export_inputs, monkeypatch):
    from app.mission.exporter.snapshot import (
        SnapshotCaptureError,
        capture_export_snapshot,
    )

    _, routes, pois, _ = export_inputs
    original = routes.get_route

    def changing(route_id):
        route = original(route_id)
        route.metadata.name += "x"
        return route

    monkeypatch.setattr(routes, "get_route", changing)
    with pytest.raises(SnapshotCaptureError, match="changed"):
        capture_export_snapshot("m", routes, pois)


def test_capture_freezes_default_coverage_before_preparation(
    export_inputs, monkeypatch, tmp_path
):
    from app.mission.exporter.snapshot import capture_export_snapshot
    from app.mission.exporter.snapshot_views import SnapshotViews

    _, routes, pois, _ = export_inputs
    path = tmp_path / "data/sat_coverage/commka.geojson"
    path.parent.mkdir(parents=True)
    path.write_text(
        json.dumps(
            {
                "type": "FeatureCollection",
                "features": [
                    {
                        "type": "Feature",
                        "properties": {"satellite_id": "AOR"},
                        "geometry": {
                            "type": "Polygon",
                            "coordinates": [
                                [[-2, -2], [2, -2], [2, 2], [-2, 2], [-2, -2]]
                            ],
                        },
                    }
                ],
            }
        )
    )
    monkeypatch.setattr("app.mission.timeline_service._COVERAGE_SAMPLER", None)
    original = timeline_preparation.prepare_mission_timeline

    def prepare(*args, **kwargs):
        path.unlink()
        assert kwargs["coverage_sampler"].check_coverage_at_point(0, 0) == ["AOR"]
        return original(*args, **kwargs)

    monkeypatch.setattr(
        "app.mission.exporter.snapshot.prepare_mission_timeline", prepare
    )
    snapshot = capture_export_snapshot("m", routes, pois)
    assert SnapshotViews(snapshot).timeline("l").segments
    assert snapshot.legs[0].preparation_origin == "rebuilt"


def test_capture_discovers_packaged_kmz_without_writing_coverage(
    export_inputs, monkeypatch, tmp_path
):
    from app.mission.exporter.snapshot import capture_export_snapshot

    _, routes, pois, _ = export_inputs
    monkeypatch.setattr("app.mission.timeline_service._COVERAGE_SAMPLER", None)
    snapshot = capture_export_snapshot("m", routes, pois)
    assert snapshot.legs[0].preparation_origin == "rebuilt"
    assert any(p.name == "coverage/kmz" and p.content for p in snapshot.source_payloads)
    assert not (tmp_path / "data/sat_coverage").exists()


def test_fallback_bounds_use_effective_splice_and_current_departure(
    export_inputs, monkeypatch
):
    from app.mission.exporter.snapshot import capture_export_snapshot

    mission, routes, pois, _ = export_inputs
    leg = mission.legs[0]
    leg.adjusted_departure_time = BASE + timedelta(hours=2)
    leg.transports.manual_aar_tracks = [
        ManualAARTrack(
            id="track",
            name="Track",
            points=[
                ManualAARTrackPoint(latitude=0.3, longitude=0.01),
                ManualAARTrackPoint(latitude=0.7, longitude=0.01),
            ],
        )
    ]
    leg.transports.manual_route_splice = ManualRouteSplice(
        enabled_track_id="track", speed_knots=100
    )
    storage.save_mission_v2(mission)
    prepared = timeline_preparation.prepare_mission_timeline(leg, routes, pois)
    storage.save_mission_timeline("l", prepared.timeline, "m")
    monkeypatch.setattr(
        "app.mission.exporter.snapshot.prepare_mission_timeline",
        Mock(side_effect=RuntimeError("failure")),
    )
    snapshot = capture_export_snapshot("m", routes, pois)
    assert snapshot.legs[0].preparation_origin == "cached"
    assert json.loads(
        snapshot.legs[0].effective_route_json
    ) == prepared.route.model_dump(mode="json")
    assert snapshot.legs[0].utc_bounds == (
        prepared.projector.start_time,
        prepared.projector.end_time,
    )


def test_splice_coincident_joins_remain_mappable_without_changing_effective_route(
    export_inputs,
):
    from app.mission.exporter.customer_document import build_customer_mission_document
    from app.mission.exporter.snapshot import capture_export_snapshot

    mission, routes, pois, _ = export_inputs
    leg = mission.legs[0]
    leg.transports.manual_aar_tracks = [
        ManualAARTrack(
            id="track",
            name="Track",
            points=[
                ManualAARTrackPoint(latitude=0.25, longitude=0),
                ManualAARTrackPoint(latitude=0.5, longitude=0.25),
                ManualAARTrackPoint(latitude=0.75, longitude=0),
            ],
        )
    ]
    leg.transports.manual_route_splice = ManualRouteSplice(
        enabled_track_id="track",
        leave_segment_index=0,
        leave_fraction=0.25,
        rejoin_segment_index=0,
        rejoin_fraction=0.75,
        speed_knots=450,
    )
    storage.save_mission_v2(mission)
    captured = capture_export_snapshot("m", routes, pois)
    original = captured.legs[0].effective_route_json
    points = json.loads(original)["points"]
    assert any(
        a["expected_arrival_time"] == b["expected_arrival_time"]
        for a, b in pairwise(points)
    )
    payload = build_customer_mission_document(captured)["legs"][0]
    assert payload["mapInput"] is not None
    scene = payload["mapInput"]["route"]
    assert len({p["timestamp"] for p in scene}) == len(scene)
    assert scene[0]["timestamp"] == payload["flight"]["startUtc"]
    assert scene[-1]["timestamp"] == payload["flight"]["endUtc"]
    assert any(p["latitude"] == 0.5 and p["longitude"] == 0.25 for p in scene)
    assert any("coincident" in reason for reason in payload["mapInputDiagnostics"])
    assert captured.legs[0].effective_route_json == original
    assert routes.get_route("r").points[-1].expected_arrival_time == BASE + timedelta(
        hours=1
    )


def test_save_can_finish_during_preparation_without_changing_capture(
    export_inputs, monkeypatch
):
    from concurrent.futures import ThreadPoolExecutor

    from app.mission.exporter.snapshot import capture_export_snapshot
    from app.mission.exporter.snapshot_views import SnapshotViews

    mission, routes, pois, _ = export_inputs
    original = timeline_preparation.prepare_mission_timeline

    def prepare(*args, **kwargs):
        def save():
            with storage.get_active_leg_lock(), storage.get_mission_lock("m"):
                mission.name = "Saved while preparing"
                storage.save_mission_v2(mission)

        with ThreadPoolExecutor(max_workers=1) as executor:
            executor.submit(save).result(timeout=2)
        return original(*args, **kwargs)

    monkeypatch.setattr(
        "app.mission.exporter.snapshot.prepare_mission_timeline", prepare
    )
    snapshot = capture_export_snapshot("m", routes, pois)
    assert SnapshotViews(snapshot).mission().name == "Captured mission"
    assert storage.load_mission_v2("m").name == "Saved while preparing"


@pytest.mark.parametrize("continuous", [False, True])
def test_capture_handles_dependency_disappearing_during_read(
    export_inputs, monkeypatch, continuous
):
    from app.mission.exporter.snapshot import (
        SnapshotCaptureError,
        capture_export_snapshot,
    )

    _, routes, pois, _ = export_inputs
    original = routes.get_route
    calls = 0

    def read(route_id):
        nonlocal calls
        calls += 1
        if continuous or calls == 1:
            raise FileNotFoundError("Concurrent dependency replacement")
        return original(route_id)

    monkeypatch.setattr(routes, "get_route", read)
    if continuous:
        with pytest.raises(SnapshotCaptureError, match="unavailable"):
            capture_export_snapshot("m", routes, pois)
    else:
        snapshot = capture_export_snapshot("m", routes, pois)
        assert snapshot.legs[0].preparation_origin == "rebuilt"
        assert any("retry" in warning.lower() for warning in snapshot.warnings)


def test_cached_fallback_keeps_committed_map_markers(export_inputs, monkeypatch):
    from app.mission.exporter.snapshot import capture_export_snapshot
    from app.mission.exporter.snapshot_views import SnapshotViews

    _, routes, pois, _ = export_inputs
    pois._pois["p"].generated_source = "mission-timeline"
    pois._pois["p"].kind = "departure"
    timeline = timeline_preparation.prepare_mission_timeline(
        storage.load_mission_v2("m").legs[0], routes, pois
    ).timeline
    storage.save_mission_timeline("l", timeline, "m")
    monkeypatch.setattr(
        "app.mission.exporter.snapshot.prepare_mission_timeline",
        Mock(side_effect=RuntimeError("failure")),
    )
    snapshot = capture_export_snapshot("m", routes, pois)
    pois._pois["p"].name = "Later marker"
    assert snapshot.legs[0].preparation_origin == "cached"
    assert (
        SnapshotViews(snapshot).poi_manager.list_pois()[0].generated_source
        == "mission-timeline"
    )
    assert [
        p.name for p in SnapshotViews(snapshot).map_poi_manager("l").list_pois()
    ] == ["Captured POI"]


def test_rebuilt_map_replaces_stale_generated_markers_without_changing_source_export(
    export_inputs,
):
    from app.mission.exporter.snapshot import capture_export_snapshot
    from app.mission.exporter.snapshot_views import SnapshotViews

    _, routes, pois, _ = export_inputs
    pois._pois["p"].generated_source = "mission-timeline"
    pois._pois["p"].kind = "departure"
    snapshot = capture_export_snapshot("m", routes, pois)
    views = SnapshotViews(snapshot)
    assert views.poi_manager.list_pois()[0].name == "Captured POI"
    assert "Captured POI" not in [
        p.name for p in views.map_poi_manager("l").list_pois()
    ]
    assert "generated_source" not in views.poi_manager.list_pois()[0].model_dump(
        mode="json"
    )


def test_capture_uninitialized_catalog_is_read_only(
    export_inputs, monkeypatch, tmp_path
):
    import app.satellites.catalog as catalog_module
    from app.mission.exporter.snapshot import capture_export_snapshot

    _, routes, pois, _ = export_inputs
    monkeypatch.setattr(catalog_module, "_catalog", None)
    snapshot = capture_export_snapshot("m", routes, pois)
    assert snapshot.legs[0].preparation_origin == "rebuilt"
    assert catalog_module._catalog is None
    assert not (tmp_path / "data/satellites").exists()
    assert not (tmp_path / "data/sat_coverage").exists()


@pytest.mark.parametrize("content", [b"", None])
def test_source_archive_distinguishes_empty_kml_from_missing_file(
    export_inputs, content
):
    from app.mission.exporter.snapshot import capture_export_snapshot

    _, routes, pois, _ = export_inputs
    path = routes.routes_dir / "r.kml"
    if content is None:
        path.unlink()
    else:
        path.write_bytes(content)
    snapshot = capture_export_snapshot("m", routes, pois)
    name = "kml_missing/r" if content is None else "kml/r"
    assert any(s.name == name and s.content == b"" for s in snapshot.source_payloads)


@pytest.mark.parametrize("blocked_first", [False, True])
def test_export_captures_material_x_changes_without_changing_legacy_events(
    export_inputs, monkeypatch, blocked_first
):
    from app.mission.exporter.customer_projection import project_briefing_leg
    from app.mission.exporter.snapshot import capture_export_snapshot
    from app.mission.models import KaOutage, KuOutageOverride
    from app.satellites.rules import EventType, RuleEngine

    mission, routes, pois, catalog = export_inputs
    mission.legs[0].transports.aar_windows = []
    mission.legs[0].transports.ka_outages = [
        KaOutage(id="ka", start_time=BASE, duration_seconds=3600)
    ]
    mission.legs[0].transports.ku_overrides = [
        KuOutageOverride(id="ku", start_time=BASE, duration_seconds=3600)
    ]
    storage.save_mission_v2(mission)

    def geometry(self, *, timestamp, **kwargs):
        blocked = (timestamp < BASE + timedelta(minutes=30)) == blocked_first
        return (
            True,
            180.0,
            {
                "violation_reason": "elevation" if blocked else "azimuth",
                "elevation_degrees": 0.0 if blocked else 35.0,
                "min_elevation_degrees": 10.0,
                "elevation_below_min": blocked,
            },
        )

    monkeypatch.setattr(RuleEngine, "evaluate_x_azimuth_window", geometry)
    legacy = timeline_preparation.prepare_mission_timeline(
        mission.legs[0], routes, pois, satellite_catalog=catalog
    )
    snap = capture_export_snapshot("m", routes, pois)
    rebuilt = json.loads(snap.legs[0].timeline_json)
    assert (
        rebuilt["segments"] == json.loads(legacy.timeline.model_dump_json())["segments"]
    )
    assert [
        e.timestamp
        for e in legacy.events
        if e.event_type == EventType.X_AZIMUTH_VIOLATION
    ] == (
        [
            BASE + timedelta(minutes=1),
            BASE + timedelta(minutes=1),
            BASE + timedelta(minutes=30),
            BASE + timedelta(hours=1),
        ]
        if blocked_first
        else [
            BASE + timedelta(minutes=1),
            BASE + timedelta(minutes=30),
            BASE + timedelta(hours=1),
            BASE + timedelta(hours=1),
        ]
    )
    projection = project_briefing_leg(snap.legs[0])
    a = next(
        i
        for i in projection.intervals
        if i.start_time <= BASE + timedelta(minutes=20) < i.end_time
    )
    b = next(
        i
        for i in projection.intervals
        if i.start_time <= BASE + timedelta(minutes=40) < i.end_time
    )
    assert [a.decisions[2].value, b.decisions[2].value] == ["Down", "Down"]
    blocked = a if blocked_first else b
    assert blocked.remaining_transports == ()
    assert blocked.posture == "Communications unavailable"
