"""Customer deck contracts, checked against editable PowerPoint objects."""

import io
import zipfile
from collections import Counter
from dataclasses import replace
from itertools import pairwise
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from app.mission.exporter.snapshot import ExportSnapshot
from app.mission.exporter.snapshot_inputs import canonical_json
from app.mission.exporter.trial_clocks import format_clocks
from app.mission.exporter.trial_maps import LegMapResult, MapView
from app.mission.exporter.trial_projection import project_trial_leg
from tests.unit.customer_briefing_fixtures import fixture, snapshot


def inputs(name="F02_AR"):
    data = fixture(name)
    captured = snapshot(data)
    trial = project_trial_leg(captured)
    image = io.BytesIO()
    Image.new("RGB", (640, 360), "#eeeeee").save(image, "PNG")
    maps = LegMapResult(
        captured.leg_id,
        (MapView("view-1", image.getvalue()),),
        "Overview map unavailable — static route fallback",
        "static",
        (),
        "test-key",
        ("Origin", "Destination"),
    )
    frozen = ExportSnapshot(
        data["mission"]["id"],
        "fixture-fingerprint",
        canonical_json(data["mission"]),
        (captured,),
        (),
        (),
    )
    return frozen, (trial,), (maps,)


def build(name="F02_AR"):
    from app.mission.exporter.trial_pptx import build_trial_pptx

    args = inputs(name)
    return Presentation(io.BytesIO(build_trial_pptx(*args))), args


def texts(slide):
    return "\n".join(
        (
            shape.text
            if shape.has_text_frame
            else (
                "\n".join(c.text for row in shape.table.rows for c in row.cells)
                if shape.has_table
                else ""
            )
        )
        for shape in slide.shapes
    )


def primary(prs):
    return [s for s in prs.slides if not s.name.startswith("appendix")]


def runs(shape):
    frames = [shape.text_frame] if shape.has_text_frame else []
    if shape.has_table:
        frames += [c.text_frame for row in shape.table.rows for c in row.cells]
    return [r for frame in frames for p in frame.paragraphs for r in p.runs if r.text]


def test_trial_slide_contract_and_editability():
    prs, (frozen, legs, _) = build()
    leg = legs[0]
    assert abs(prs.slide_width / Inches(1) - 13.333) < 0.001
    assert prs.slide_height == Inches(7.5)
    timeline_slides = [
        s for s in primary(prs) if any(sh.name.startswith("bar:") for sh in s.shapes)
    ]
    covered = []
    for slide in timeline_slides:
        labels = [sh for sh in slide.shapes if sh.name.startswith("lane:")]
        assert [" ".join(sh.text.split()) for sh in labels] == [
            "Overall communications posture",
            "Commercial Ka",
            "Starshield",
            "X-Band MILSATCOM",
            "SOF / AR restrictions",
        ]
        assert labels[0].height >= labels[1].height * 2
        assert [sh.top for sh in labels] == sorted(sh.top for sh in labels)
        bars = [sh for sh in slide.shapes if sh.name.startswith("bar:")]
        for interval in leg.intervals:
            matching = [sh for sh in bars if sh.name == f"bar:{interval.id}:0"]
            if matching:
                covered.append(interval.id)
        for bar in bars:
            lane = int(bar.name.rsplit(":", 1)[1])
            if lane:
                xml = bar._element.xml
                assert all(
                    color not in xml
                    for color in ("39834A", "E4B83F", "DD8736", "BD4040")
                )
    assert Counter(covered) == Counter(i.id for i in leg.intervals)
    tables = [sh.table for s in primary(prs) for sh in s.shapes if sh.has_table]
    assert tables
    headers = [
        "Start (ET)",
        "End (ET)",
        "Event / impact",
        "Communications remaining",
        "Overall posture / customer implication",
    ]
    assert all(
        [" ".join(c.text.split()) for c in t.rows[0].cells] == headers for t in tables
    )
    table_text = "\n".join(
        c.text for t in tables for row in list(t.rows)[1:] for c in row.cells
    )
    for row in leg.coordination_rows:
        assert row.id in table_text
        assert (
            format_clocks(row.start_time, leg.utc_bounds[0]).et.replace(" ", "\n", 1)
            in table_text
        )
        assert (
            format_clocks(row.end_time, leg.utc_bounds[0]).et.replace(" ", "\n", 1)
            in table_text
        )
        for source_id in row.active_source_ids:
            all_primary = "\n".join(texts(s) for s in primary(prs))
            assert (
                source_id in all_primary.replace("\n", "") or source_id in all_primary
            )
    for slide in prs.slides:
        for shape in slide.shapes:
            for run in runs(shape):
                minimum = (
                    14
                    if slide.name.startswith("appendix")
                    or shape.name.startswith("support:")
                    else 18
                )
                if shape.name.startswith("time:") or shape.name.startswith("window:"):
                    minimum = 20
                assert run.font.size is not None and run.font.size.pt >= minimum
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left + shape.width <= prs.slide_width + 5
            assert shape.top + shape.height <= prs.slide_height + 5
    logo = next(sh for sh in prs.slides[0].shapes if sh.name == "apo-patch")
    with Image.open(Path("app/mission/assets/APO Patch.jpg")) as image:
        assert abs(logo.width / logo.height - image.width / image.height) < 0.001
    assert any(
        "pattFill" in sh._element.xml for s in timeline_slides for sh in s.shapes
    )
    result = io.BytesIO()
    prs.save(result)
    with zipfile.ZipFile(result) as z:
        assert len([n for n in z.namelist() if n.startswith("ppt/media/")]) >= 2
        assert not any(
            b'TargetMode="External"' in z.read(n)
            for n in z.namelist()
            if n.endswith(".rels")
        )
    tables[0].cell(1, 2).text = "Edited offline"
    restored = io.BytesIO()
    prs.save(restored)
    assert any(
        "Edited offline" in texts(s)
        for s in Presentation(io.BytesIO(restored.getvalue())).slides
    )
    assert frozen.fingerprint == "fixture-fingerprint"


def test_trial_dense_pagination():
    prs, (_, legs, _) = build("F09")
    leg = legs[0]
    pages = primary(prs)
    assert len(pages) > 10
    assert fixture("F09")["route"]["points"][0]["name"] in " ".join(
        texts(pages[0]).split()
    )
    assert fixture("F09")["route"]["points"][-1]["name"] in " ".join(
        texts(pages[0]).split()
    )
    content = "\n".join(texts(s) for s in pages)
    for interval in leg.coordination_rows:
        assert interval.id in content.replace("\n", "") or interval.id in content
        facts = []
        for slide in pages:
            if slide.name.startswith("window-detail") and any(
                sh.name == f"window:{interval.id}" for sh in slide.shapes
            ):
                facts.extend(
                    sh.text for sh in slide.shapes if sh.name == "window-facts"
                )
        combined = " ".join(" ".join(facts).split())
        assert all(" ".join(cause.split()) in combined for cause in interval.causes)
    for slide in pages:
        text = texts(slide)
        assert "Leg 1 of 1" in text
        assert "Up" in text and "Down" in text and "?" in text and "SOF / AR" in text
        assert "EDT" in text and "Z" in text and "T+" in text
        assert "prediction, not a throughput guarantee" in text
    assert "continued" in content.lower()


def test_trial_source_appendix():
    prs, (_, legs, _) = build("F02_AR")
    source_pages = [s for s in prs.slides if s.name.startswith("appendix-source")]
    text = "\n".join(texts(s) for s in source_pages)
    for source in legs[0].sources:
        assert text.count("Source: " + source.source_id + "\n") == 1
        assert source.reason in text.replace("\n", " ")
        for timestamp in (source.start_time, source.end_time, source.timestamp):
            if timestamp:
                clock = format_clocks(timestamp, legs[0].utc_bounds[0])
                assert (
                    clock.et in text and clock.zulu in text and clock.relative in text
                )
        assert source.source_revision in text
        assert source.source_digest in text.replace("\n", "")
    appendix = "\n".join(texts(s) for s in prs.slides if s.name.startswith("appendix"))
    for interval in legs[0].intervals:
        assert interval.id in appendix
        assert all(
            format_clocks(t, legs[0].utc_bounds[0]).relative in appendix
            for t in (interval.start_time, interval.end_time)
        )


@pytest.mark.parametrize("name", ["F01", "F03", "F04", "F06", "F10"])
def test_quiet_brief_uncertain_and_missing_contracts(name):
    prs, (_, legs, _) = build(name)
    text = "\n".join(texts(s) for s in primary(prs))
    if name == "F01":
        assert legs[0].quiet_summary in text.replace("\n", " ")
        assert all(i.id in text for i in legs[0].coordination_rows)
    if name == "F03":
        assert ":30" in text
        assert "Communications unavailable" in text.replace("\n", " ")
        assert any(
            sh.name.startswith("callout:") for s in primary(prs) for sh in s.shapes
        )
    if name == "F06":
        assert "Posture uncertain" in " ".join(text.split())
        table_postures = " ".join(
            c.text
            for s in primary(prs)
            for sh in s.shapes
            if sh.has_table
            for row in list(sh.table.rows)[1:]
            for c in [row.cells[4]]
        )
        assert "Communications unavailable" not in " ".join(table_postures.split())
        assert "Up*" not in text
    if name == "F10":
        assert "incomplete" in text.lower() or "unavailable" in text.lower()
        assert "No communications degradation" not in text


def test_multi_leg_offline_views_and_snapshot_unchanged():
    from app.mission.exporter.trial_pptx import build_trial_pptx

    template, _, maps = inputs("F01")
    data = fixture("F08")["legs"]
    captured = tuple(snapshot(item) for item in data)
    legs = tuple(project_trial_leg(item) for item in captured)
    maps = tuple(
        replace(
            maps[0],
            leg_id=item.leg_id,
            views=(
                maps[0].views[0],
                replace(maps[0].views[0], id="view-2"),
                replace(maps[0].views[0], id="view-3"),
            ),
        )
        for item in captured
    )
    frozen = replace(template, legs=captured)
    original = repr(frozen)
    prs = Presentation(io.BytesIO(build_trial_pptx(frozen, legs, maps)))
    assert repr(frozen) == original
    content = "\n".join(texts(s) for s in prs.slides)
    assert all(f"Leg {n} of 3" in content for n in range(1, 4))
    assert all(day in content for day in ("2026-10-07", "2026-10-10", "2026-10-13"))
    assert "view-3" in content
    for leg in legs:
        timeline = next(
            s for s in prs.slides if s.name.startswith(f"timeline:{leg.leg_id}:")
        )
        assert "T+00:00" in texts(timeline)


@pytest.mark.parametrize(
    "fault",
    ["leg-identity", "map-identity", "invalid-image", "partition", "filtered-boundary"],
)
def test_trial_generation_failure_is_atomic(fault):
    from app.mission.exporter.trial_pptx import TrialGenerationError, build_trial_pptx

    frozen, legs, maps = inputs()
    if fault == "leg-identity":
        legs = (replace(legs[0], leg_id="different"),)
    if fault == "map-identity":
        maps = (replace(maps[0], leg_id="different"),)
    if fault == "invalid-image":
        maps = (replace(maps[0], views=(MapView("bad", b"not PNG"),)),)
    if fault == "partition":
        legs = (replace(legs[0], intervals=legs[0].intervals[1:]),)
    if fault == "filtered-boundary":
        legs = (
            replace(
                legs[0],
                coordination_rows=(
                    replace(
                        legs[0].coordination_rows[0], end_time=legs[0].utc_bounds[1]
                    ),
                ),
            ),
        )
    with pytest.raises(TrialGenerationError):
        build_trial_pptx(frozen, legs, maps)


def test_unmapped_warnings_and_unavailable_map_card():
    from app.mission.exporter.trial_pptx import build_trial_pptx

    frozen, legs, maps = inputs("F01")
    note = "Window 3: route position unavailable"
    maps = (
        replace(
            maps[0],
            views=(),
            status="unavailable",
            label="Route map unavailable",
            reason="Route data missing",
            warnings=(note,),
        ),
    )
    prs = Presentation(io.BytesIO(build_trial_pptx(frozen, legs, maps)))
    text = " ".join("\n".join(texts(s) for s in primary(prs)).split())
    assert "Route map unavailable" in text
    assert note in text


def test_untimed_leg_retains_unsplit_sources_without_inventing_tzero():
    from app.mission.exporter.trial_pptx import build_trial_pptx

    frozen, _, maps = inputs("F02_AR")
    captured = replace(frozen.legs[0], utc_bounds=None)
    frozen = replace(frozen, legs=(captured,))
    leg = project_trial_leg(captured)
    prs = Presentation(io.BytesIO(build_trial_pptx(frozen, (leg,), maps)))
    appendix = " ".join(
        "\n".join(texts(s) for s in prs.slides if s.name.startswith("appendix")).split()
    )
    assert all(source.source_id in appendix for source in leg.sources)
    assert "T-zero unavailable" in appendix
    assert "T+00:00" not in appendix


def test_primary_window_and_clock_labels_at_twenty_points():
    prs, _ = build("F03")
    for slide in primary(prs):
        for shape in slide.shapes:
            if shape.has_table:
                for row in list(shape.table.rows)[1:]:
                    assert all(
                        r.font.size.pt >= 20
                        for c in list(row.cells)[:2]
                        for p in c.text_frame.paragraphs
                        for r in p.runs
                        if r.text
                    )
                    first = row.cells[2].text_frame.paragraphs[0]
                    assert first.text.startswith("Window ")
                    assert all(r.font.size.pt >= 20 for r in first.runs if r.text)


def test_timeline_bars_remain_proportional_and_panels_touch():
    prs, (_, legs, _) = build("F03")
    leg = legs[0]
    seen = []
    for slide in primary(prs):
        bars = [
            sh
            for sh in slide.shapes
            if sh.name.startswith("bar:") and sh.name.endswith(":0")
        ]
        if not bars:
            continue
        intervals = [
            next(i for i in leg.intervals if sh.name == f"bar:{i.id}:0") for sh in bars
        ]
        scale = sum(sh.width for sh in bars) / sum(
            (i.end_time - i.start_time).total_seconds() for i in intervals
        )
        for bar, interval in zip(bars, intervals):
            assert (
                abs(
                    bar.width
                    - (interval.end_time - interval.start_time).total_seconds() * scale
                )
                < 5
            )
        assert all(abs(a.left + a.width - b.left) < 5 for a, b in pairwise(bars))
        seen.extend(intervals)
    assert tuple(seen) == leg.intervals


def test_quiet_rows_share_page_and_long_reasons_expand_as_linked_details():
    prs, _ = build("F01")
    assert any(
        sh.has_table and len(sh.table.rows) == 3
        for s in primary(prs)
        for sh in s.shapes
    )
    dense, _ = build("F09")
    assert any(s.name.startswith("window-detail") for s in primary(dense))
    assert len(primary(dense)) < 100


def test_short_coordination_summaries_are_not_needlessly_split():
    prs, (_, legs, _) = build("F02")
    actual = Counter()
    for slide in primary(prs):
        for shape in slide.shapes:
            if shape.has_table:
                for row in list(shape.table.rows)[1:]:
                    label = row.cells[2].text.replace("\n", "")
                    actual.update(
                        i.id for i in legs[0].coordination_rows if i.id in label
                    )
    assert actual == Counter(i.id for i in legs[0].coordination_rows)


def test_validation_rejects_a_serialized_font_contract_violation(monkeypatch):
    from pptx.presentation import Presentation as PresentationClass

    from app.mission.exporter.trial_pptx import TrialGenerationError, build_trial_pptx

    original_save = PresentationClass.save

    def corrupt(prs, destination):
        title = next(sh for sh in prs.slides[0].shapes if sh.has_text_frame)
        title.text_frame.paragraphs[0].runs[0].font.size = Inches(12 / 72)
        original_save(prs, destination)

    monkeypatch.setattr(PresentationClass, "save", corrupt)
    with pytest.raises(TrialGenerationError):
        build_trial_pptx(*inputs("F01"))


def test_grayscale_posture_labels_and_body_legend_remain_explicit():
    prs, (_, legs, _) = build("F06")
    uncertain = {i.id for i in legs[0].intervals if i.posture == "Posture uncertain"}
    for slide in primary(prs):
        content = " ".join(texts(slide).split())
        assert (
            "Communications unavailable" in content and "Posture uncertain" in content
        )
        for shape in slide.shapes:
            if shape.name in {f"bar:{identity}:0" for identity in uncertain}:
                assert shape.text in ("?", "Posture uncertain")
            if shape.name in ("legend", "prediction-caveat"):
                assert all(run.font.size.pt >= 18 for run in runs(shape))
