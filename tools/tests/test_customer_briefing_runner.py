"""Acceptance must reject mutable inputs and reap owned timeout workers."""

import importlib.util
import json
import os
import subprocess
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[2]
RUNNER = ROOT / "tools/acceptance/customer-briefing/generate.py"


def load():
    spec = importlib.util.spec_from_file_location("briefing_acceptance", RUNNER)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


@pytest.fixture
def checkout(tmp_path):
    root = tmp_path / "repo"
    root.mkdir()
    (root / "input").write_text("candidate")
    for args in (
        ["init", "-q"],
        ["add", "."],
        [
            "-c",
            "user.name=Test",
            "-c",
            "user.email=test@example.com",
            "commit",
            "-qm",
            "candidate",
        ],
    ):
        subprocess.run(["git", *args], cwd=root, check=True)
    sha = subprocess.check_output(
        ["git", "rev-parse", "HEAD"], cwd=root, text=True
    ).strip()
    return root, sha


def test_requires_full_matching_clean_candidate(checkout):
    runner = load()
    root, sha = checkout
    runner.validate_candidate(root, sha)
    for invalid in (sha[:8], "0" * 40):
        with pytest.raises(ValueError):
            runner.validate_candidate(root, invalid)
    (root / "input").write_text("uncommitted")
    with pytest.raises(ValueError, match="clean"):
        runner.validate_candidate(root, sha)


def test_refuses_existing_task_root_without_removing_it(tmp_path):
    runner = load()
    task = tmp_path / "task"
    task.mkdir()
    marker = task / "shared"
    marker.write_text("preserve")
    with pytest.raises(ValueError, match="exists"):
        runner.claim_paths(tmp_path / "evidence", task)
    assert marker.read_text() == "preserve"


def test_timeout_reaps_real_child_group_and_retains_ownership(tmp_path):
    runner = load()
    owner = runner.Owner(tmp_path)
    with pytest.raises(subprocess.TimeoutExpired):
        owner.run(
            "hung",
            [
                sys.executable,
                "-c",
                'import subprocess,time; subprocess.Popen(["sleep","120"]); time.sleep(120)',
            ],
            seconds=0.2,
        )
    info = json.loads((tmp_path / "hung-ownership.json").read_text())
    assert info["pid"] and info["pgid"] == info["pid"]
    with pytest.raises(ProcessLookupError):
        os.killpg(info["pgid"], 0)
    assert info["group_gone"]


def test_failed_command_preserves_log_and_fails_gate(tmp_path):
    runner = load()
    with pytest.raises(subprocess.CalledProcessError):
        runner.Owner(tmp_path).run(
            "failed",
            [sys.executable, "-c", 'print("failure evidence"); raise SystemExit(7)'],
        )
    assert "failure evidence" in (tmp_path / "failed.log").read_text()
    assert (
        json.loads((tmp_path / "failed-ownership.json").read_text())["exit_code"] == 7
    )


def test_signal_reaps_owned_group(tmp_path):
    script = tmp_path / "interrupt_owner.py"
    script.write_text(
        f'import importlib.util,signal\ns=importlib.util.spec_from_file_location("runner",{str(RUNNER)!r});m=importlib.util.module_from_spec(s);s.loader.exec_module(m)\nsignal.signal(signal.SIGTERM, lambda *a: (_ for _ in ()).throw(SystemExit(143)))\nm.Owner({str(tmp_path)!r}).run("signal-child",["sleep","120"])\n'
    )
    process = subprocess.Popen([sys.executable, str(script)])
    import time

    try:
        deadline = time.monotonic() + 5
        record = tmp_path / "signal-child-ownership.json"
        while not record.exists() or not json.loads(record.read_text()).get("pid"):
            assert time.monotonic() < deadline
            time.sleep(0.02)
        process.terminate()
        process.wait(timeout=12)
        info = json.loads(record.read_text())
        with pytest.raises(ProcessLookupError):
            os.killpg(info["pgid"], 0)
    finally:
        if process.poll() is None:
            process.kill()
        process.wait()


def test_pptx_inspector_rejects_external_media(tmp_path):
    import zipfile

    inspect_path = ROOT / "tools/acceptance/customer-briefing/inspect_pptx.py"
    spec = importlib.util.spec_from_file_location("inspect_briefing", inspect_path)
    inspector = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(inspector)
    deck = tmp_path / "external.pptx"
    with zipfile.ZipFile(deck, "w") as archive:
        archive.writestr(
            "ppt/slides/_rels/slide1.xml.rels",
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Target="https://private/image.png" TargetMode="External" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"/></Relationships>',
        )
    with pytest.raises(ValueError, match="external"):
        inspector.check_archive(deck.read_bytes())


def test_fixture_matrix_covers_dst_and_all_three_f08_legs():
    runner = load()
    fixtures = runner.fixture_cases(ROOT)
    assert {
        "F01",
        "F02",
        "F02_AR",
        "F03",
        "F04",
        "F05",
        "F06",
        "F07_midnight",
        "F07_spring",
        "F07_fall",
        "F08",
        "F09",
        "F10",
    } == set(fixtures)
    assert len(fixtures["F08"]) == 3
    assert fixtures["F07_spring"][0]["utc_bounds"] == [
        "2026-03-08T06:50:00Z",
        "2026-03-08T07:20:00Z",
    ]
    assert fixtures["F07_fall"][0]["utc_bounds"] == [
        "2026-11-01T05:50:00Z",
        "2026-11-01T06:20:00Z",
    ]


def test_project_collision_never_removes_existing_images(
    checkout, tmp_path, monkeypatch
):
    from types import SimpleNamespace

    runner = load()
    root, sha = checkout
    monkeypatch.setattr(runner, "ROOT", root)
    monkeypatch.setattr(runner, "port_free", lambda port: None)
    calls = []

    def fake_run(self, name, command, **kwargs):
        calls.append(command)
        # Existing Compose resources are external to this attempt.
        text = "existing-container\n" if name == "preflight-containers" else ""
        if name.startswith("image-check"):
            text = "existing-image\n"
        (self.evidence / f"{name}.log").write_text(text)

    monkeypatch.setattr(runner.Owner, "run", fake_run)
    args = SimpleNamespace(
        sha=sha,
        port=15309,
        evidence=tmp_path / "evidence",
        task_root=tmp_path / "task",
        check=False,
        profile=None,
        api_only=True,
    )
    with pytest.raises((ValueError, RuntimeError)):
        runner.production(args)
    assert not any(command[:3] == ["docker", "image", "rm"] for command in calls)
    assert not any("down" in command for command in calls)


def test_import_archive_retains_route_ids_sources_and_dst(checkout, tmp_path):
    import zipfile

    runner = load()
    output = tmp_path / "imports"
    imports = runner.import_archives(ROOT, output)
    with zipfile.ZipFile(imports["F08"]["archive"]) as archive:
        mission = json.loads(archive.read("mission.json"))
        assert len(mission["legs"]) == 3
        assert len({leg["route_id"] for leg in mission["legs"]}) == 3
        for leg in mission["legs"]:
            assert archive.read(f'routes/{leg["route_id"]}.kml').startswith(b"<?xml")
            assert (
                json.loads(archive.read(f'pois/{leg["id"]}-pois.json'))["pois"][0][
                    "name"
                ]
                == "Synthetic retained source POI"
            )
    with zipfile.ZipFile(imports["F07_fall"]["archive"]) as archive:
        assert b"2026-11-01 05:50:00Z" in archive.read("routes/f07-fall-route.kml")
        assert b"2026-11-01 06:20:00Z" in archive.read("routes/f07-fall-route.kml")


def test_log_collection_failure_still_tears_down_project(
    checkout, tmp_path, monkeypatch
):
    from types import SimpleNamespace

    runner = load()
    root, sha = checkout
    monkeypatch.setattr(runner, "ROOT", root)
    monkeypatch.setattr(runner, "port_free", lambda port: None)
    monkeypatch.setattr(runner, "http", lambda *a, **k: (b"", {}))
    monkeypatch.setattr(
        runner,
        "import_archives",
        lambda *a: (_ for _ in ()).throw(RuntimeError("stop after startup")),
    )
    calls = []

    def fake_run(self, name, command, **kwargs):
        calls.append(command)
        text = (
            json.dumps(
                [{"Config": {"Labels": {"org.opencontainers.image.revision": sha}}}]
            )
            if name == "image-identities"
            else ""
        )
        (self.evidence / f"{name}.log").write_text(text)
        if name == "container-logs":
            raise subprocess.CalledProcessError(1, command)

    monkeypatch.setattr(runner.Owner, "run", fake_run)
    args = SimpleNamespace(
        sha=sha,
        port=15309,
        evidence=tmp_path / "evidence",
        task_root=tmp_path / "task",
        check=False,
        profile=None,
        api_only=True,
    )
    with pytest.raises(RuntimeError):
        runner.production(args)
    assert any("down" in command and "--volumes" in command for command in calls)
    assert not args.task_root.exists()


def test_backend_probe_exports_the_real_application(tmp_path, monkeypatch):
    import pathlib
    from types import ModuleType, SimpleNamespace

    app = object()
    monkeypatch.setattr(pathlib.Path, "mkdir", lambda *a, **k: None)
    main = ModuleType("main")
    main.app = app
    package = ModuleType("app.mission.package")
    package.__main__ = SimpleNamespace(capture_export_snapshot=lambda *a: None)
    exporter = ModuleType("app.mission.exporter")
    exporter.trial_maps = SimpleNamespace(render_trial_maps=lambda *a: None)
    exporter.trial_pptx = SimpleNamespace(build_trial_pptx=lambda *a: None)
    exporter.trial_projection = SimpleNamespace(project_trial_leg=lambda *a: None)
    for name, module in [
        ("main", main),
        ("app", ModuleType("app")),
        ("app.mission", ModuleType("app.mission")),
        ("app.mission.package", package),
        ("app.mission.exporter", exporter),
    ]:
        monkeypatch.setitem(sys.modules, name, module)
    spec = importlib.util.spec_from_file_location(
        "backend_probe", ROOT / "tools/acceptance/customer-briefing/backend_probe.py"
    )
    probe = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(probe)
    assert getattr(probe, "app", None) is app


def test_existing_package_anonymization_preserves_times_and_references(tmp_path):
    import zipfile

    spec = importlib.util.spec_from_file_location(
        "anonymize", ROOT / "tools/acceptance/customer-briefing/anonymize_package.py"
    )
    helper = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(helper)
    original = tmp_path / "private.zip"
    with zipfile.ZipFile(original, "w") as archive:
        archive.writestr(
            "mission.json",
            json.dumps(
                {
                    "id": "private-mission",
                    "name": "Private customer",
                    "description": "Private crew",
                    "metadata": {"mission_number": "PRIVATE"},
                    "legs": [
                        {
                            "id": "private-leg",
                            "name": "Private leg",
                            "route_id": "private-route",
                            "adjusted_departure_time": "2026-10-07T08:00:00Z",
                            "transports": {
                                "aar_windows": [
                                    {
                                        "id": "private-ar",
                                        "start_waypoint_name": "Private waypoint",
                                        "end_waypoint_name": "Private waypoint",
                                    }
                                ]
                            },
                        }
                    ],
                }
            ),
        )
        archive.writestr(
            "routes/private-route.kml",
            '<kml xmlns="http://www.opengis.net/kml/2.2"><Document><Placemark><name>Private waypoint</name><description>Private crew; Time Over Waypoint: 2026-10-07 08:00:00Z</description><Point><coordinates>1,2,0</coordinates></Point></Placemark></Document></kml>',
        )
    out = tmp_path / "anonymized.zip"
    report = helper.anonymize(original, out)
    with zipfile.ZipFile(out) as archive:
        assert all(
            b"Private" not in archive.read(n) and b"private-" not in archive.read(n)
            for n in archive.namelist()
        )
        mission = json.loads(archive.read("mission.json"))
        leg = mission["legs"][0]
        route = archive.read(f'routes/{leg["route_id"]}.kml')
        assert (
            leg["transports"]["aar_windows"][0]["start_waypoint_name"].encode() in route
        )
        assert b"Time Over Waypoint: 2026-10-07 08:00:00Z" in route
        assert b"1,2,0" in route
        assert leg["adjusted_departure_time"] == "2026-10-07T08:00:00Z"
    assert report["geometry"] == "retained locally; no customer data committed"


def test_source_comparison_accepts_legacy_serialization_but_rejects_time_change():
    runner = load()
    api = {
        "initial_x_satellite_id": "X-fixture",
        "ka_outages": [
            {
                "id": "outage",
                "start_time": "2026-10-07T08:10:00Z",
                "duration_seconds": 30,
                "reason": None,
            }
        ],
    }
    legacy = {
        "initial_x_satellite_id": "X-fixture",
        "ka_outages": [
            {
                "id": "outage",
                "start_time": "2026-10-07 08:10:00+00:00",
                "duration_seconds": 30,
            }
        ],
    }
    assert runner.source_transport_equal(api, legacy)
    legacy["ka_outages"][0]["start_time"] = "2026-10-07 08:10:01+00:00"
    assert not runner.source_transport_equal(api, legacy)


def test_port_check_accepts_time_wait_and_rejects_live_listener():
    import socket

    runner = load()
    with socket.socket() as listener:
        listener.setsockopt(socket.SOL_SOCKET, socket.SO_REUSEADDR, 1)
        listener.bind(("127.0.0.1", 0))
        port = listener.getsockname()[1]
        listener.listen()
        with pytest.raises(OSError):
            runner.port_free(port)
        with socket.socket() as client:
            client.connect(("127.0.0.1", port))
            peer, _ = listener.accept()
            peer.close()
    # The listener is gone; a closed accepted connection is not a live resource.
    runner.port_free(port)


@pytest.mark.parametrize("reordered", [False, True])
@pytest.mark.parametrize("alter_poi", [False, True])
def test_existing_exports_preserve_actual_pois_and_match_legs_by_identity(
    tmp_path, monkeypatch, alter_poi, reordered
):
    import io
    import zipfile

    from pptx import Presentation

    runner = load()
    monkeypatch.syspath_prepend(str(RUNNER.parent))
    legs = [
        {
            "id": str(i),
            "route_id": "r" + str(i),
            "transports": {"initial_x_satellite_id": "X-" + str(i)},
        }
        for i in (1, 2)
    ]
    original = tmp_path / "source.zip"
    deck = io.BytesIO()
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(deck)
    exported = io.BytesIO()
    with zipfile.ZipFile(original, "w") as source, zipfile.ZipFile(
        exported, "w"
    ) as result:
        result.writestr("manifest.json", "{}")
        result.writestr("exports/mission/mission-slides.pptx", deck.getvalue())
        for leg in legs:
            result.writestr(f'legs/{leg["id"]}.json', json.dumps(leg))
            poi = {
                "name": "Anonymized marker",
                "latitude": 1.0,
                "longitude": 2.0,
                "category": "mission-event",
            }
            path = f'pois/{leg["id"]}-pois.json'
            source.writestr(path, json.dumps({"pois": [poi]}))
            result.writestr(
                path,
                json.dumps({"pois": [{**poi, "longitude": 3.0 if alter_poi else 2.0}]}),
            )
    before = json.dumps({"legs": list(reversed(legs)) if reordered else legs}).encode()

    def response(base, path, **kwargs):
        if path.endswith("/export"):
            return exported.getvalue(), {
                "X-Mission-Export-Trial-Status": "disabled",
                "X-Mission-Export-Warnings": "[]",
            }
        return before, {}

    monkeypatch.setattr(runner, "http", response)
    args = (
        "http://private",
        {
            "EXISTING": {
                "mission_id": "m",
                "mission": {"legs": legs},
                "archive": str(original),
                "required_poi_name": None,
            }
        },
        tmp_path / "outputs",
        "disabled",
    )
    if alter_poi:
        with pytest.raises(AssertionError, match="POI"):
            runner.api_exports(*args)
    else:
        assert runner.api_exports(*args)["EXISTING"]["source_preservation"] == "pass"


def test_anonymization_keeps_satellite_lookup_targets_consistent(tmp_path):
    import zipfile

    spec = importlib.util.spec_from_file_location(
        "anonymize", RUNNER.parent / "anonymize_package.py"
    )
    helper = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(helper)
    source, output = tmp_path / "source.zip", tmp_path / "out.zip"
    with zipfile.ZipFile(source, "w") as z:
        z.writestr(
            "mission.json",
            json.dumps(
                {
                    "id": "m",
                    "legs": [
                        {
                            "id": "l",
                            "transports": {
                                "initial_x_satellite_id": "X-1",
                                "x_transitions": [
                                    {"target_satellite_id": "Private Satellite"}
                                ],
                            },
                        }
                    ],
                }
            ),
        )
        z.writestr(
            "pois/satellites.json",
            json.dumps(
                {
                    "pois": [
                        {"id": "p1", "name": "X-1", "longitude": 12},
                        {"id": "p2", "name": "Private Satellite", "longitude": 34},
                    ]
                }
            ),
        )
    helper.anonymize(source, output)
    with zipfile.ZipFile(output) as z:
        transports = json.loads(z.read("mission.json"))["legs"][0]["transports"]
        pois = {
            p["name"]: p["longitude"]
            for p in json.loads(z.read("pois/satellites.json"))["pois"]
        }
    assert transports["initial_x_satellite_id"] == "X-1"
    assert pois[transports["initial_x_satellite_id"]] == 12
    assert pois[transports["x_transitions"][0]["target_satellite_id"]] == 34
    assert "Private Satellite" not in pois


def test_nginx_export_timeout_is_scoped_and_exceeds_trial_and_legacy_budget():
    import re

    config = (ROOT / "frontend/mission-planner/nginx.conf").read_text()
    blocks = re.findall(r"location\s+([^{}]+)\{([^{}]+)\}", config)
    export = [(location, body) for location, body in blocks if "/export" in location]
    assert len(export) == 1
    location, body = export[0]
    assert re.search(r"/api/v2/missions/.+/export", location)
    assert "proxy_pass http://starlink-location:8000;" in body
    seconds = int(re.search(r"proxy_read_timeout (\d+)s;", body).group(1))
    assert seconds >= 180
    assert all(
        "proxy_read_timeout" not in body
        for location, body in blocks
        if "/export" not in location
    )


@pytest.mark.parametrize("unsafe", [False, True])
def test_container_evidence_stream_copies_real_archive_without_daemon_cp(
    tmp_path, monkeypatch, unsafe
):
    import io
    import tarfile

    runner = load()
    owner = runner.Owner(tmp_path / "evidence")

    def streamed(self, name, command, **kwargs):
        assert "exec" in command and "cp" not in command
        buffer = io.BytesIO()
        with tarfile.open(fileobj=buffer, mode="w") as archive:
            info = tarfile.TarInfo("../escape" if unsafe else "canonical/trial.pptx")
            value = b"retained offline deck"
            info.size = len(value)
            archive.addfile(info, io.BytesIO(value))
        (self.evidence / (name + ".log")).write_bytes(buffer.getvalue())

    monkeypatch.setattr(runner.Owner, "run", streamed)
    if unsafe:
        with pytest.raises(ValueError, match="unsafe"):
            runner.copy_probe(
                owner,
                ["docker", "compose"],
                {},
                "copy",
                "canonical",
                tmp_path / "output",
            )
        assert not (tmp_path / "escape").exists()
    else:
        runner.copy_probe(
            owner, ["docker", "compose"], {}, "copy", "canonical", tmp_path / "output"
        )
        assert (tmp_path / "output/trial.pptx").read_bytes() == b"retained offline deck"
