#!/usr/bin/env python3
"""Task 4 builder-only previews. Does not export ZIPs or launch production maps.

Run under a recorded process owner and GNU timeout. Outputs contain synthetic
fixtures only; pair/ZIP/UI acceptance remains a separate Task 6 deliverable.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT / "backend/starlink-location"))


def generate(evidence):
    from pptx import Presentation

    from app.mission.exporter.snapshot import ExportSnapshot
    from app.mission.exporter.snapshot_inputs import canonical_json
    from app.mission.exporter.trial_maps import render_trial_maps
    from app.mission.exporter.trial_pptx import build_trial_pptx
    from app.mission.exporter.trial_projection import project_trial_leg
    from tests.unit.customer_briefing_fixtures import fixture, snapshot

    evidence.mkdir(parents=True, exist_ok=True)
    manifest = {}
    for name in ("F01", "F02_AR", "F06", "F09"):
        data = fixture(name)
        captured = snapshot(data)
        frozen = ExportSnapshot(
            data["mission"]["id"],
            hashlib.sha256(canonical_json(data)).hexdigest(),
            canonical_json(data["mission"]),
            (captured,),
            (),
            (),
        )
        legs = (project_trial_leg(captured),)
        # Exercise final IDs in the existing bounded neutral fallback. The
        # production scene remains qualified by Task 3B, not by this preview.
        maps = render_trial_maps(frozen, legs, budget_seconds=0)
        deck = build_trial_pptx(frozen, legs, maps)
        (evidence / f"{name}.pptx").write_bytes(deck)
        prs = Presentation(io.BytesIO(deck))
        manifest[name] = {
            "snapshot_fingerprint": frozen.fingerprint,
            "deck_sha256": hashlib.sha256(deck).hexdigest(),
            "slides": [
                {"number": n, "name": slide.name}
                for n, slide in enumerate(prs.slides, 1)
            ],
            "primary_windows": [
                {
                    "id": i.id,
                    "number": i.window_number,
                    "start": i.start_time.isoformat(),
                    "end": i.end_time.isoformat(),
                    "posture": i.posture,
                }
                for i in legs[0].coordination_rows
            ],
            "map_status": maps[0].status,
            "map_warnings": maps[0].warnings,
        }
        print(f"{name}: {len(prs.slides)} slides, {maps[0].status} map", flush=True)
    (evidence / "decks.json").write_text(json.dumps(manifest, indent=2) + "\n")


def sheets(evidence):
    from PIL import Image, ImageDraw

    manifest = json.loads((evidence / "decks.json").read_text())
    for name, info in manifest.items():
        pages = sorted((evidence / name).glob("page-*.png"))
        assert len(pages) == len(info["slides"]), (
            name,
            len(pages),
            len(info["slides"]),
        )
        primary = [
            pages[s["number"] - 1]
            for s in info["slides"]
            if not s["name"].startswith("appendix")
        ]
        for page in primary:
            with Image.open(page) as image:
                image.convert("L").save(page.with_name(page.stem + "-gray.png"))
        for offset in range(0, len(primary), 12):
            batch = primary[offset : offset + 12]
            canvas = Image.new("RGB", (1200, 4 * 258), "#eeeeee")
            draw = ImageDraw.Draw(canvas)
            for index, page in enumerate(batch):
                x, y = (index % 3) * 400, (index // 3) * 258
                with Image.open(page) as image:
                    image.thumbnail((398, 224))
                    canvas.paste(image, (x, y + 24))
                draw.text((x + 8, y + 5), f"{name} {page.stem}", fill="black")
            canvas.save(evidence / f"{name}-primary-{offset // 12 + 1}.png")
        info["rendered_pages"] = len(pages)
        info["grayscale_primary_pages"] = len(primary)
    (evidence / "decks.json").write_text(json.dumps(manifest, indent=2) + "\n")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("action", choices=("generate", "sheets"))
    parser.add_argument("--evidence", type=Path, required=True)
    args = parser.parse_args()
    {"generate": generate, "sheets": sheets}[args.action](args.evidence.resolve())
