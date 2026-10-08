#!/usr/bin/env python3
"""Inspect actual PPTX objects and offline asset relationships; retain slide links."""

from __future__ import annotations

import hashlib
import io
import json
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET


def check_archive(data):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        if archive.testzip():
            raise ValueError("corrupt PPTX archive")
        for name in archive.namelist():
            if name.endswith(".rels"):
                for rel in ET.fromstring(archive.read(name)):
                    if rel.attrib.get(
                        "TargetMode"
                    ) == "External" and not rel.attrib.get("Type", "").endswith(
                        "/hyperlink"
                    ):
                        raise ValueError(f"external asset: {name}")
        return {
            "parts": len(archive.namelist()),
            "media": {
                n: hashlib.sha256(archive.read(n)).hexdigest()
                for n in archive.namelist()
                if n.startswith("ppt/media/")
            },
        }


def inspect(path, trial=False):
    from pptx import Presentation

    report = check_archive(path.read_bytes())
    prs = Presentation(path)
    report.update(
        sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
        slides=[],
        width_inches=round(prs.slide_width / 914400, 3),
        height_inches=round(prs.slide_height / 914400, 3),
    )
    for number, slide in enumerate(prs.slides, 1):
        fonts, texts, tables = [], [], 0
        for shape in slide.shapes:
            if shape.has_table:
                tables += 1
                frames = [
                    cell.text_frame for row in shape.table.rows for cell in row.cells
                ]
            else:
                frames = [shape.text_frame] if shape.has_text_frame else []
            for frame in frames:
                texts.append(frame.text)
                for para in frame.paragraphs:
                    for run in para.runs:
                        if run.text.strip() and run.font.size:
                            fonts.append(run.font.size.pt)
            if (
                trial
                and shape.left >= 0
                and (
                    shape.left + shape.width > prs.slide_width + 9144
                    or shape.top + shape.height > prs.slide_height + 9144
                )
            ):
                raise ValueError(f"shape outside slide {number}: {shape.name}")
        report["slides"].append(
            {
                "number": number,
                "name": slide.name,
                "text": texts,
                "editable_tables": tables,
                "font_sizes_pt": sorted(set(fonts)),
                "primary": not slide.name.startswith(("appendix", "mission-index")),
            }
        )
    if not report["slides"]:
        raise ValueError("empty deck")
    if trial:
        sys.path.insert(
            0, str(Path(__file__).resolve().parents[3] / "backend/starlink-location")
        )
        from app.mission.exporter.trial_pptx import validate_trial_pptx

        validate_trial_pptx(path.read_bytes())
    return report


if __name__ == "__main__":
    deck = Path(sys.argv[1])
    print(json.dumps(inspect(deck, "--trial" in sys.argv), indent=2))
